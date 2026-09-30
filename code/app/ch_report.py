"""Kanallarim: haftalik hisobot (PDF + Excel). Dushanba 09:00 da avtomatik (dastur ochiq bo'lsa) yoki qo'lda."""
import asyncio
import json
from datetime import datetime, timedelta

from . import ai, ch_data, ch_extra, ch_insight, ch_stats, ch_track, db, notify
from .config import DATA_DIR, log

FMT = "%Y-%m-%d %H:%M:%S"
REPORT_DIR = DATA_DIR / "reports" / "ch"


def _sum(rows, key):
    return sum((r[key] or 0) for r in rows)


def gather(cid: int, days: int = 7) -> dict:
    ch = ch_data.get(cid)
    now = datetime.now()
    cur_from = (now - timedelta(days=days)).strftime("%Y-%m-%d")
    prev_from = (now - timedelta(days=2 * days)).strftime("%Y-%m-%d")
    d = ch_insight.daily(cid, 2 * days)
    cur, prev = d[days:], d[:days]
    mem_now = next((r["members"] for r in reversed(d) if r["members"] is not None), ch["members"])
    mem_start = next((r["members"] for r in cur if r["members"] is not None), None)
    top = ch_stats.top_posts(cid, 0, days, 5, by="views")
    links = [t for t in ch_track.top_links(cid, 10) if t["n"]]
    abs_ = [a for a in ch_extra.ab_list(cid) if a.get("eval") and a["eval"].get("ready")][:3]
    alerts = db.q("SELECT title, ts FROM ch_alerts WHERE channel_id=? AND ts>=? ORDER BY id DESC LIMIT 10", (cid, cur_from))
    sent = db.one("SELECT COUNT(*) c FROM ch_plan WHERE channel_id=? AND status='sent' AND sent_at>=?", (cid, cur_from))["c"]
    failed = db.one("SELECT COUNT(*) c FROM ch_plan WHERE channel_id=? AND status IN ('failed','missed') AND created_at>=?", (cid, prev_from))["c"]
    usage = db.one("SELECT COALESCE(SUM(cost),0) c, COALESCE(SUM(tokens_in+tokens_out),0) t FROM ch_usage WHERE channel_id=? AND day>=?", (cid, cur_from))
    aud = ch_insight.audience_latest(cid)
    rep, rep_at = ch_extra.last_comments_report(cid)
    return {"channel": ch["title"], "days": days, "from": cur_from, "to": now.strftime("%Y-%m-%d"),
            "members": mem_now, "net": (mem_now - mem_start) if (mem_now is not None and mem_start is not None) else None,
            "posts": _sum(cur, "posts"), "posts_prev": _sum(prev, "posts"), "views": _sum(cur, "views"), "views_prev": _sum(prev, "views"),
            "avg": round(_sum(cur, "views") / _sum(cur, "posts")) if _sum(cur, "posts") else 0,
            "clicks": _sum(cur, "clicks"), "clicks_prev": _sum(prev, "clicks"),
            "joined": _sum([r for r in cur if r["joined"] is not None], "joined") if any(r["joined"] is not None for r in cur) else None,
            "left": _sum([r for r in cur if r["left"] is not None], "left") if any(r["left"] is not None for r in cur) else None,
            "daily": cur, "top_posts": [{"date": p["date"][:16], "media": p["media"], "views": p["views"], "err": p["err"],
                                          "text": (p["text"] or "")[:120].replace("\n", " ")} for p in top],
            "links": links, "ab": [{"title": a["title"], "verdict": a["eval"]["verdict"], "winner": a["winner"]} for a in abs_],
            "alerts": [dict(a) for a in alerts], "plan_sent": sent, "plan_failed": failed,
            "ai_cost": float(usage["c"] or 0), "ai_tokens": int(usage["t"] or 0),
            "audience": {k: aud[k] for k in ("active_pct", "premium_pct", "coverage", "sampled", "total")} if aud else None,
            "comments_summary": (rep or {}).get("summary") if rep else None, "comments_at": rep_at}


SUMMARY_SYSTEM = """You write the weekly summary for a business Telegram channel owner (goal: sales and leads) from real numbers only.
Write in Uzbek (Latin), plain and concrete; compare with the previous week; do not invent numbers.
Return ONLY JSON: {"summary": "4-6 sentences", "wins": ["..."], "risks": ["..."], "actions": ["3-5 concrete next steps for next week"]}"""


async def ai_summary(cid: int, data: dict) -> dict | None:
    if not ai.configured():
        return None
    slim = {k: v for k, v in data.items() if k not in ("daily",)}
    try:
        res, _ = await ai.chat(cid, "analysis", SUMMARY_SYSTEM, json.dumps(slim, ensure_ascii=False, default=str)[:9000],
                               max_out=1800, note="haftalik hisobot")
        return res if isinstance(res, dict) else None
    except ai.AIError as e:
        log.warning("Haftalik xulosa olinmadi: %s", e)
        return None


def _pct(a, b):
    if not b:
        return "—"
    return f"{(a - b) / b * 100:+.0f}%"


def build_xlsx(data: dict, summ: dict | None, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    head = PatternFill("solid", fgColor="0F766E")
    ws.append([f"{data['channel']} — haftalik hisobot", f"{data['from']} … {data['to']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Ko'rsatkich", "Shu hafta", "Oldingi hafta", "O'zgarish"])
    for c in ws[3]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    ws.append(["Obunachilar (hozir)", data["members"], "", data["net"]])
    ws.append(["Postlar", data["posts"], data["posts_prev"], _pct(data["posts"], data["posts_prev"])])
    ws.append(["Ko'rishlar (shu hafta chiqqan postlar)", data["views"], data["views_prev"], _pct(data["views"], data["views_prev"])])
    ws.append(["Havola bosishlari", data["clicks"], data["clicks_prev"], _pct(data["clicks"], data["clicks_prev"])])
    if data["joined"] is not None:
        ws.append(["Qo'shildi / chiqib ketdi (rasmiy)", f"{data['joined']} / {data['left']}", "", ""])
    ws.append(["AI xarajati (USD)", round(data["ai_cost"], 3), "", ""])
    if summ:
        ws.append([])
        ws.append(["AI xulosa"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        ws.append([summ.get("summary", "")])
        for k, t in (("wins", "Yutuqlar"), ("risks", "Xavflar"), ("actions", "Keyingi hafta")):
            for x in summ.get(k) or []:
                ws.append([t, x])
    ws.column_dimensions["A"].width = 44
    for col in "BCD":
        ws.column_dimensions[col].width = 20
    w2 = wb.create_sheet("Kunlar")
    w2.append(["Sana", "Obunachi", "O'zgarish", "Postlar", "Ko'rishlar", "O'rtacha", "Bosishlar"])
    for c in w2[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for r in data["daily"]:
        w2.append([r["day"], r["members"], r["net"], r["posts"], r["views"], r["avg"], r["clicks"]])
    w3 = wb.create_sheet("Eng yaxshi postlar")
    w3.append(["Sana", "Tur", "Ko'rish", "ER %", "Matn"])
    for c in w3[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for p in data["top_posts"]:
        w3.append([p["date"], p["media"], p["views"], p["err"], p["text"]])
    w4 = wb.create_sheet("Havolalar")
    w4.append(["Post", "Kod", "Variant", "Bosishlar"])
    for c in w4[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for t in data["links"]:
        w4.append([t["label"], t["code"], t["variant"], t["n"]])
    wb.save(path)


def build_pdf(data: dict, summ: dict | None, path, colors_hex=("#0F766E", "#F59E0B", "#111827")):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .reports import _font
    font = _font()
    main = colors.HexColor(colors_hex[0])
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName=font, textColor=main)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=font, textColor=main)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName=font, leading=14)

    def esc(t):
        return str(t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def table(rows, widths=None):
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 9),
                               ("BACKGROUND", (0, 0), (-1, 0), main), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        return t
    el = [Paragraph(f"{esc(data['channel'])}: haftalik hisobot", h1), Paragraph(f"{data['from']} … {data['to']}", body), Spacer(1, 5 * mm)]
    el.append(table([["Ko'rsatkich", "Shu hafta", "Oldingi", "O'zgarish"],
                     ["Obunachilar", data["members"] if data["members"] is not None else "—", "", f"{data['net']:+d}" if data["net"] is not None else "—"],
                     ["Postlar", data["posts"], data["posts_prev"], _pct(data["posts"], data["posts_prev"])],
                     ["Ko'rishlar", data["views"], data["views_prev"], _pct(data["views"], data["views_prev"])],
                     ["Havola bosishlari", data["clicks"], data["clicks_prev"], _pct(data["clicks"], data["clicks_prev"])],
                     ["AI xarajati (USD)", f"{data['ai_cost']:.3f}", "", ""]], [70 * mm, 35 * mm, 35 * mm, 30 * mm]))
    if summ:
        el += [Spacer(1, 5 * mm), Paragraph("AI xulosa", h2), Paragraph(esc(summ.get("summary")), body)]
        for k, t in (("wins", "Yutuqlar"), ("risks", "Xavflar"), ("actions", "Keyingi hafta uchun")):
            if summ.get(k):
                el.append(Paragraph(f"<b>{t}</b>", body))
                el += [Paragraph("• " + esc(x), body) for x in summ[k]]
    if data["top_posts"]:
        el += [Spacer(1, 5 * mm), Paragraph("Eng ko'p ko'rilgan postlar", h2),
               table([["Sana", "Tur", "Ko'rish", "ER %", "Matn"]] +
                     [[p["date"], p["media"], p["views"], p["err"], Paragraph(esc(p["text"][:90]), body)] for p in data["top_posts"]],
                     [30 * mm, 20 * mm, 20 * mm, 16 * mm, 88 * mm])]
    if data["links"]:
        el += [Spacer(1, 5 * mm), Paragraph("Kuzatuv havolalari (bosishlar)", h2),
               table([["Post", "Variant", "Bosish"]] + [[Paragraph(esc(t["label"] or t["code"]), body), t["variant"] or "", t["n"]] for t in data["links"]],
                     [110 * mm, 30 * mm, 30 * mm])]
    if data["ab"]:
        el += [Spacer(1, 5 * mm), Paragraph("A/B sinovlar", h2)] + [Paragraph(f"• {esc(a['title'])}: {esc(a['verdict'])}", body) for a in data["ab"]]
    if data["alerts"]:
        el += [Spacer(1, 5 * mm), Paragraph("Viral signallar", h2)] + [Paragraph("• " + esc(a["title"]), body) for a in data["alerts"]]
    if data.get("comments_summary"):
        el += [Spacer(1, 5 * mm), Paragraph("Auditoriya izohlari", h2), Paragraph(esc(data["comments_summary"]), body)]
    doc.build(el)


async def make(cid: int, days: int = 7, with_ai: bool = True) -> int:
    ch = ch_data.get(cid)
    data = gather(cid, days)
    summ = await ai_summary(cid, data) if with_ai else None
    out = REPORT_DIR / str(cid)
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d")
    xl, pdf = out / f"hisobot_{stamp}.xlsx", out / f"hisobot_{stamp}.pdf"
    build_xlsx(data, summ, xl)
    build_pdf(data, summ, pdf, ch_data.colors(ch))
    period = f"{data['from']}…{data['to']}"
    rid = db.ex("INSERT INTO ch_reports(channel_id,kind,period,path,created_at) VALUES(?,?,?,?,?)", (cid, "pdf", period, str(pdf), db.now()))
    db.ex("INSERT INTO ch_reports(channel_id,kind,period,path,created_at) VALUES(?,?,?,?,?)", (cid, "xlsx", period, str(xl), db.now()))
    return rid


async def loop():
    await asyncio.sleep(420)
    while True:
        try:
            now = datetime.now()
            if now.weekday() == 0 and now.hour >= 9 and db.get_setting("ch_weekly_on", "1") != "0":
                key = now.strftime("%Y-%m-%d")
                for ch in ch_data.channels():
                    if db.get_setting(f"ch_rep_{ch['id']}") == key:
                        continue
                    await make(ch["id"])
                    db.set_setting(f"ch_rep_{ch['id']}", key)
                    notify.toast("Haftalik hisobot tayyor", ch["title"])
                    notify.event("info", "hisobot", f"{ch['title']}: haftalik hisobot yaratildi")
        except asyncio.CancelledError:
            raise
        except Exception:
            log.exception("ch_report.loop")
            notify.event("error", "hisobot", "Haftalik hisobotni yaratishda xato (loglarga qarang)")
        await asyncio.sleep(1800)
