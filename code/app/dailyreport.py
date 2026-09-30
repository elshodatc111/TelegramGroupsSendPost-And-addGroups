"""Kunlik hisobot (Excel): qo'lda yuklab olish va har kuni avtomatik saqlash."""
import asyncio
import io
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import db
from .config import DATA_DIR, log
from .limits import get_warmup, health

REPORT_DIR = DATA_DIR / "reports"
RED = "DC2626"


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RED)
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def build(aid: int, day: str) -> bytes:
    a = db.one("SELECT * FROM accounts WHERE id=?", (aid,))
    like = day + "%"
    posts = db.q("SELECT jt.sent_at, jt.title, jt.status, jt.error, jt.views, jt.reactions, jt.replies, j.name, jv.idx "
                 "FROM job_targets jt JOIN jobs j ON j.id=jt.job_id LEFT JOIN job_variants jv ON jv.id=jt.variant_id "
                 "WHERE j.account_id=? AND jt.sent_at LIKE ? ORDER BY jt.sent_at", (aid, like))
    joins = db.q("SELECT jt.tried_at, jt.ref, jt.title, jt.status, jt.detail FROM join_targets jt "
                 "JOIN join_batches b ON b.id=jt.batch_id WHERE b.account_id=? AND jt.tried_at LIKE ? ORDER BY jt.tried_at",
                 (aid, like))
    left = db.q("SELECT ts, title, reason, auto FROM leave_log WHERE account_id=? AND ts LIKE ? ORDER BY ts", (aid, like))
    sent = [p for p in posts if p["status"] == "sent"]
    failed = [p for p in posts if p["status"] == "failed"]
    ok_join = [j for j in joins if j["status"] in ("joined", "already", "requested")]
    h, w = health(aid), get_warmup(aid)
    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    _sheet(ws, ["Ko'rsatkich", "Qiymat"], [
        ["Akkaunt", a["name"] if a else aid], ["Sana", day],
        ["Yuborilgan postlar", len(sent)], ["Xato bilan", len(failed)],
        ["Jami ko'rishlar", sum(p["views"] or 0 for p in sent)], ["Reaksiyalar", sum(p["reactions"] or 0 for p in sent)],
        ["Javoblar", sum(p["replies"] or 0 for p in sent)],
        ["A'zo bo'lish urinishlari", len(joins)], ["Muvaffaqiyatli a'zo bo'lish", len(ok_join)],
        ["Guruhdan chiqildi", len(left)], ["Akkaunt salomatligi", f"{h['label']} ({h['score']})"],
        ["Isitish rejasi", "Faol" if w and w["active"] else "Yo'q"]], [32, 30])
    _sheet(wb.create_sheet("Postlar"), ["Vaqt", "Guruh", "Kampaniya", "Variant", "Holat", "Ko'rish", "Reaksiya", "Javob", "Izoh"],
           [[p["sent_at"], p["title"], p["name"] or "", (p["idx"] or 0) + 1, "Yuborildi" if p["status"] == "sent" else "Xato",
             p["views"], p["reactions"], p["replies"], p["error"] or ""] for p in posts], [19, 34, 22, 9, 12, 10, 10, 10, 40])
    _sheet(wb.create_sheet("A'zo bo'lish"), ["Vaqt", "Manba", "Guruh", "Holat", "Izoh"],
           [[j["tried_at"], j["ref"], j["title"] or "", j["status"], j["detail"] or ""] for j in joins], [19, 30, 32, 14, 40])
    _sheet(wb.create_sheet("Chiqishlar"), ["Vaqt", "Guruh", "Sabab", "Kim"],
           [[l["ts"], l["title"], l["reason"], "Avtomatik" if l["auto"] else "Qo'lda"] for l in left], [19, 34, 44, 12])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save(aid: int, day: str):
    d = REPORT_DIR / str(aid)
    d.mkdir(parents=True, exist_ok=True)
    (d / f"kunlik_{day}.xlsx").write_bytes(build(aid, day))


def files(aid: int) -> list[str]:
    d = REPORT_DIR / str(aid)
    return sorted((f.name for f in d.glob("kunlik_*.xlsx")), reverse=True)[:60] if d.exists() else []


async def loop():
    """Har kuni tunda kechagi hisobotni saqlab qo'yadi."""
    await asyncio.sleep(120)
    while True:
        try:
            yday = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
            for a in db.q("SELECT id FROM accounts WHERE workspace='posting'"):
                if not (REPORT_DIR / str(a["id"]) / f"kunlik_{yday}.xlsx").exists():
                    save(a["id"], yday)
        except asyncio.CancelledError:
            raise
        except Exception:
            log.exception("dailyreport")
        await asyncio.sleep(1800)
