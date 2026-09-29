"""Excel/CSV dan Telegram username va havolalarni o'qish, hisobotni Excel'ga eksport qilish."""
import csv
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

USER = r"[A-Za-z][A-Za-z0-9_]{3,31}"
INVITE_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:t|telegram)\.me/(?:\+|joinchat/)([A-Za-z0-9_-]{8,})", re.I)
LINK_RE = re.compile(rf"(?:https?://)?(?:www\.)?(?:t|telegram)\.me/({USER})(?:[/?#].*)?$", re.I)
AT_RE = re.compile(rf"^@({USER})$")
PLAIN_RE = re.compile(rf"^({USER})$")
NOT_NAMES = {"joinchat", "addlist", "share", "username", "usernames", "channel", "channels", "kanal",
             "kanallar", "group", "groups", "guruh", "guruhlar", "link", "links", "havola", "name",
             "title", "telegram", "nomi", "user", "users", "url", "urls"}

SUPPORTED = (".xlsx", ".xlsm", ".csv", ".txt")


def parse_cell(v):
    """Katak matnidan (kind, key, ref) ni ajratadi yoki None."""
    if v is None:
        return None
    s = str(v).strip()
    if not s or len(s) > 200:
        return None
    m = INVITE_RE.search(s)
    if m:
        return "invite", m.group(1), s
    m = LINK_RE.match(s)
    if m and m.group(1).lower() not in NOT_NAMES:
        return "username", m.group(1), s
    m = AT_RE.match(s) or PLAIN_RE.match(s)
    if m and m.group(1).lower() not in NOT_NAMES:
        return "username", m.group(1), s
    return None


def _read_sheets(path: Path) -> dict[str, list[list]]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        wb = load_workbook(path, data_only=True)
        sheets = {}
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows():
                out = []
                for cell in row:
                    v = cell.value
                    link = getattr(cell.hyperlink, "target", None) if cell.hyperlink else None
                    if link and parse_cell(v) is None and parse_cell(link):
                        v = link      # katakda oddiy nom, havola esa yashirin bo'lsa
                    out.append(v)
                rows.append(out)
            sheets[ws.title] = rows
        return sheets
    if ext in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if ext == ".txt":
            return {"Matn": [[ln] for ln in text.splitlines()]}
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return {"CSV": list(csv.reader(io.StringIO(text), dialect))}
    raise ValueError("Faqat .xlsx, .csv yoki .txt fayllar qo'llab-quvvatlanadi (eski .xls ni .xlsx qilib saqlang)")


def analyze(path: Path) -> list[dict]:
    """Fayldagi har bir ustun bo'yicha topilgan username/havolalar ro'yxati."""
    columns = []
    for sheet, rows in _read_sheets(path).items():
        width = max((len(r) for r in rows), default=0)
        for ci in range(width):
            items, seen, header = [], set(), ""
            for ri, row in enumerate(rows):
                v = row[ci] if ci < len(row) else None
                p = parse_cell(v)
                if ri == 0 and p is None and v not in (None, ""):
                    header = str(v).strip()[:40]
                if not p:
                    continue
                kind, key, ref = p
                k = (kind, key.lower() if kind == "username" else key)
                if k in seen:
                    continue
                seen.add(k)
                items.append({"kind": kind, "key": key, "ref": ref[:120]})
            if items:
                label = f"{sheet} · {get_column_letter(ci + 1)}" + (f" ({header})" if header else "")
                columns.append({"label": label, "items": items})
    return columns


def export_report(batch, targets, labels: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    head = ["#", "Username / havola", "Holat", "Izoh", "Guruh nomi", "Vaqt"]
    ws.append(head)
    fill = PatternFill("solid", fgColor="2B7BF0")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center")
    colors = {"joined": "D9F3E6", "already": "D9F3E6", "requested": "FFF1C9", "pending": "EEF1F6"}
    for i, t in enumerate(targets, 1):
        ws.append([i, t["ref"], labels.get(t["status"], t["status"]), t["detail"] or "", t["title"] or "",
                   t["tried_at"] or ""])
        row = ws[ws.max_row]
        f = PatternFill("solid", fgColor=colors.get(t["status"], "FBE0E0"))
        row[2].fill = f
    for col, w in zip("ABCDEF", (6, 38, 20, 55, 36, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    sm = wb.create_sheet("Xulosa")
    counts = {}
    for t in targets:
        counts[t["status"]] = counts.get(t["status"], 0) + 1
    sm.append(["Fayl", batch["filename"]])
    sm.append(["Yaratilgan", batch["created_at"]])
    sm.append(["Jami", len(targets)])
    for k, v in counts.items():
        sm.append([labels.get(k, k), v])
    sm.column_dimensions["A"].width = 26
    sm.column_dimensions["B"].width = 30
    for r in sm.iter_rows(min_col=1, max_col=1):
        r[0].font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
