"""Telegram SMM: g'oyalar, AI sarfi (token) va sozlamalar."""
import io
import json
import os
import subprocess
import sys
from datetime import datetime

from fastapi import APIRouter, Form, Request
from fastapi.responses import Response

from . import ai, ch_ai, ch_data, ch_stats, ch_track, db, transcribe
from .config import log
from .web import go, need_channel, page

router = APIRouter()


# ---------------------------------------------------------------- g'oyalar
@router.get("/ch/ideas")
async def ch_ideas(request: Request):
    if (r := need_channel(request)):
        return r
    ch = request.state.channel
    rows = db.q("SELECT * FROM ch_ideas WHERE channel_id=? AND kind IN ('manual','daily') ORDER BY id DESC LIMIT 60", (ch["id"],))
    mem = ch_data.ensure_memory(ch["id"])
    return page(request, "ch_ideas.html", ch=ch, ideas=rows, ai_ok=ai.configured(), mem=mem,
                n_comp=len(ch_data.competitors(ch["id"])), month_cost=ai.month_cost(ch["id"]),
                default_lang={"ru": "ru", "mixed": "both"}.get(ch["lang"] or "uz", "uz"), idea_model=ai.model_for("idea"))


@router.post("/ch/ideas/new")
async def ch_idea_new(request: Request, lang: str = Form(""), video: str = Form("auto"), note: str = Form("")):
    ch = request.state.channel
    if not ch:
        return go("/ch/channels", err="Kanal tanlanmagan")
    try:
        iid = await ch_ai.generate_idea(ch["id"], lang, video, note)
    except ai.AIError as e:
        return go("/ch/ideas", err=str(e))
    except Exception as e:
        log.exception("G'oya yaratishda xato")
        return go("/ch/ideas", err=f"Kutilmagan xato: {e}")
    return go(f"/ch/ideas/{iid}", msg="G'oya tayyor")


def _idea(request, iid):
    ch = request.state.channel
    row = db.one("SELECT * FROM ch_ideas WHERE id=?", (iid,))
    if not ch or not row or row["channel_id"] != ch["id"]:
        return None, None, None
    return ch, row, ch_ai.idea_body(row)


@router.get("/ch/ideas/{iid}")
async def ch_idea_view(request: Request, iid: int):
    ch, row, body = _idea(request, iid)
    if not row:
        return go("/ch/ideas", err="G'oya topilmadi")
    plan = db.one("SELECT * FROM ch_plan WHERE id=?", (row["plan_id"],)) if row["plan_id"] else None
    return page(request, "ch_idea.html", ch=ch, idea=row, b=body, plan=plan, now=datetime.now().strftime("%Y-%m-%dT%H:%M"),
                footer=ch_data.lead_footer(ch))


@router.post("/ch/ideas/{iid}/status")
async def ch_idea_status(request: Request, iid: int, status: str = Form(...), feedback: str = Form("")):
    ch, row, _ = _idea(request, iid)
    if not row or status not in ("new", "accepted", "rejected"):
        return go("/ch/ideas", err="G'oya topilmadi")
    db.ex("UPDATE ch_ideas SET status=?, feedback=? WHERE id=?", (status, feedback.strip()[:1000], iid))
    return go(f"/ch/ideas/{iid}", msg="Belgilandi: keyingi g'oyalar shunga qarab yaxshilanadi")


@router.post("/ch/ideas/{iid}/plan")
async def ch_idea_plan(request: Request, iid: int):
    ch, row, body = _idea(request, iid)
    if not row:
        return go("/ch/ideas", err="G'oya topilmadi")
    form = await request.form()
    lang = form.get("lang") or (body.get("langs") or ["uz"])[0]
    text = (form.get("text") or (body.get("post") or {}).get(lang) or "").strip()
    if not text:
        return go(f"/ch/ideas/{iid}", err="Post matni bo'sh")
    names, mtype = await ch_data.save_uploads(ch["id"], form.getlist("files"))
    when = (form.get("when") or "").replace("T", " ")
    if when and len(when) == 16:
        when += ":00"
    status = "scheduled" if when else "draft"
    pid = db.ex("INSERT INTO ch_plan(channel_id,title,text,media_json,media_type,parse_mode,scheduled_at,status,idea_id,created_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?)",
                (ch["id"], row["title"][:250], text, json.dumps(names), mtype, "none", when or None, status, iid, db.now()))
    ch_track.ensure_plan_code(pid, ch["id"], row["title"], idea_id=iid)
    db.ex("UPDATE ch_ideas SET status='accepted', plan_id=? WHERE id=?", (pid, iid))
    return go("/ch/plan", msg="Kontent rejaga qo'shildi" + (" (belgilangan vaqtda yuboriladi)" if when else " (qoralama)"))


@router.post("/ch/ideas/{iid}/delete")
async def ch_idea_delete(request: Request, iid: int):
    ch, row, _ = _idea(request, iid)
    if row:
        db.ex("DELETE FROM ch_ideas WHERE id=?", (iid,))
    return go("/ch/ideas", msg="G'oya o'chirildi")


# ---------------------------------------------------------------- AI sarfi
def _month(request):
    m = request.query_params.get("month") or datetime.now().strftime("%Y-%m")
    return m if len(m) == 7 and m[4] == "-" else datetime.now().strftime("%Y-%m")


@router.get("/ch/usage")
async def ch_usage(request: Request):
    month = _month(request)
    summary = ch_stats.usage_summary(month)
    sel = request.state.channel
    detail = []
    if sel:
        detail = db.q("SELECT * FROM ch_usage WHERE channel_id=? AND day LIKE ? ORDER BY id DESC LIMIT 200", (sel["id"], month + "%"))
    months = [r["m"] for r in db.q("SELECT DISTINCT SUBSTR(day,1,7) m FROM ch_usage ORDER BY m DESC LIMIT 12")]
    if month not in months:
        months.insert(0, month)
    tot = {"tin": sum(c["tin"] for c in summary), "tout": sum(c["tout"] for c in summary),
           "cost": sum(c["cost"] for c in summary), "calls": sum(c["calls"] for c in summary)}
    return page(request, "ch_usage.html", summary=summary, month=month, months=months, detail=detail, tot=tot,
                sel=None if request.state.ch_all else sel)


@router.get("/ch/usage.xlsx")
async def ch_usage_xlsx(request: Request):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    month = _month(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kanal bo'yicha"
    ws.append(["Kanal", "Oy", "So'rovlar", "Kirish token", "Chiqish token", "Audio (daqiqa)", "Xarajat (USD)", "Narx kiritilmagan"])
    for c in ch_stats.usage_summary(month):
        ws.append([c["title"], month, c["calls"], c["tin"], c["tout"], round(c["audio"] / 60, 1), round(c["cost"], 4), "ha" if c["unpriced"] else ""])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws2 = wb.create_sheet("Batafsil")
    ws2.append(["Vaqt", "Kanal", "Maqsad", "Model", "Kirish token", "Chiqish token", "Keshlangan", "Audio (soniya)", "Xarajat (USD)", "Izoh"])
    for r in db.q("SELECT * FROM ch_usage WHERE day LIKE ? ORDER BY id", (month + "%",)):
        ws2.append([r["ts"], r["ch_title"], r["purpose"], r["model"], r["tokens_in"], r["tokens_out"], r["cached_in"],
                    r["audio_sec"], r["cost"], r["note"]])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="ai_sarf_{month}.xlsx"'})


# ---------------------------------------------------------------- sozlamalar
@router.get("/ch/settings")
async def ch_settings(request: Request):
    models, models_err = [], None
    if request.query_params.get("models"):
        try:
            models = await ai.list_models()
        except ai.AIError as e:
            models_err = str(e)
    key = ai.api_key()
    return page(request, "ch_settings.html", has_key=bool(key), key_tail=(key[-4:] if key else ""),
                idea_model=ai.model_for("idea"), analysis_model=ai.model_for("analysis"), models=models, models_err=models_err,
                prices=db.q("SELECT * FROM ch_prices ORDER BY model"), cap=db.get_setting("ai_cap_usd", 0),
                learn_on=db.get_setting("ch_learn_on", "1") != "0", sync_hours=db.get_setting("ch_sync_hours", 6),
                grace=db.get_setting("ch_plan_grace_h", 3), stt_mode=db.get_setting("stt_mode") or "auto",
                whisper_model=db.get_setting("whisper_model") or "large-v3", stt_openai=db.get_setting("stt_openai_model") or "whisper-1",
                stt=transcribe.local_status(), is_win=sys.platform == "win32",
                api_id=db.get_setting("api_id", ""), has_hash=bool(db.get_setting("api_hash")),
                cf_url=ch_track.base(), has_cf_key=bool(ch_track.key()), cf_last=db.get_setting("cf_last_pull", ""))


@router.post("/ch/settings/openai")
async def ch_set_openai(key: str = Form("")):
    key = key.strip()
    if not key:
        db.del_setting("openai_key")
        return go("/ch/settings", msg="OpenAI kaliti o'chirildi")
    if not key.startswith("sk-"):
        return go("/ch/settings", err="OpenAI kaliti 'sk-' bilan boshlanadi")
    db.set_setting("openai_key", key)
    try:
        ids = await ai.list_models()
    except ai.AIError as e:
        return go("/ch/settings", err=f"Kalit saqlandi, lekin tekshiruv muvaffaqiyatsiz: {e}")
    return go("/ch/settings?models=1", msg=f"Kalit saqlandi va tekshirildi ({len(ids)} ta model mavjud). Quyida modellarni tanlang")


@router.post("/ch/settings/models")
async def ch_set_models(idea_model: str = Form(...), analysis_model: str = Form(...), cap: float = Form(0),
                        learn_on: str = Form(""), sync_hours: int = Form(6), grace: int = Form(3)):
    db.set_setting("ai_model_idea", idea_model.strip())
    db.set_setting("ai_model_analysis", analysis_model.strip())
    db.set_setting("ai_cap_usd", max(0.0, cap))
    db.set_setting("ch_learn_on", "1" if learn_on else "0")
    db.set_setting("ch_sync_hours", max(1, min(sync_hours, 48)))
    db.set_setting("ch_plan_grace_h", max(0, min(grace, 48)))
    return go("/ch/settings", msg="Saqlandi")


@router.post("/ch/settings/stt")
async def ch_set_stt(stt_mode: str = Form("auto"), whisper_model: str = Form("large-v3"), stt_openai: str = Form("whisper-1")):
    if stt_mode not in ("auto", "local", "openai", "off"):
        stt_mode = "auto"
    db.set_setting("stt_mode", stt_mode)
    db.set_setting("whisper_model", whisper_model.strip() or "large-v3")
    db.set_setting("stt_openai_model", stt_openai.strip() or "whisper-1")
    return go("/ch/settings", msg="Transkripsiya sozlamalari saqlandi")


@router.post("/ch/settings/install-whisper")
async def ch_install_whisper():
    p = transcribe.installer_path()
    if sys.platform != "win32" or not p.exists():
        return go("/ch/settings", err="O'rnatuvchi faqat Windows'da ishlaydi (tools\\install_whisper.bat)")
    subprocess.Popen(["cmd", "/c", "start", "Whisper o'rnatish", "cmd", "/k", str(p)], creationflags=0x00000010)
    return go("/ch/settings", msg="O'rnatish oynasi ochildi. Tugagach dasturni qayta ishga tushiring")


@router.post("/ch/settings/prices")
async def ch_set_prices(request: Request):
    form = await request.form()
    for m in form.getlist("model"):
        i = form.get(f"in_{m}") or "0"
        o = form.get(f"out_{m}") or "0"
        c = form.get(f"cached_{m}") or "0"
        a = form.get(f"audio_{m}") or "0"
        try:
            db.ex("UPDATE ch_prices SET in_per_m=?, out_per_m=?, cached_per_m=?, audio_per_min=? WHERE model=?",
                  (float(i), float(o), float(c), float(a), m))
        except ValueError:
            return go("/ch/settings", err=f"{m}: narx raqam bo'lishi kerak")
    new = (form.get("new_model") or "").strip()
    if new:
        try:
            db.ex("INSERT OR REPLACE INTO ch_prices(model,in_per_m,out_per_m,cached_per_m,audio_per_min) VALUES(?,?,?,?,?)",
                  (new, float(form.get("new_in") or 0), float(form.get("new_out") or 0), float(form.get("new_cached") or 0),
                   float(form.get("new_audio") or 0)))
        except ValueError:
            return go("/ch/settings", err="Yangi model narxi raqam bo'lishi kerak")
    for m in form.getlist("del"):
        db.ex("DELETE FROM ch_prices WHERE model=?", (m,))
    return go("/ch/settings", msg="Narxlar saqlandi (yangi so'rovlarga qo'llanadi)")


@router.post("/ch/settings/api")
async def ch_set_api(api_id: str = Form(...), api_hash: str = Form("")):
    api_id = api_id.strip()
    if not api_id.isdigit():
        return go("/ch/settings", err="API ID faqat raqamlardan iborat")
    db.set_setting("api_id", api_id)
    if api_hash.strip():
        db.set_setting("api_hash", api_hash.strip())
    if not db.get_setting("api_hash"):
        return go("/ch/settings", err="API HASH ni kiriting")
    return go("/ch/settings", msg="Telegram API ma'lumotlari saqlandi")


@router.post("/ch/settings/worker")
async def ch_set_worker(url: str = Form(""), key: str = Form("")):
    url = url.strip().rstrip("/")
    if not url:
        db.del_setting("cf_worker_url")
        db.del_setting("cf_key")
        return go("/ch/settings", msg="Kuzatuv havolasi o'chirildi (postlar asl havola bilan yuboriladi)")
    if not url.startswith("https://"):
        return go("/ch/settings", err="Worker manzili https:// bilan boshlanishi kerak")
    db.set_setting("cf_worker_url", url)
    if key.strip():
        db.set_setting("cf_key", key.strip())
    if not ch_track.key():
        return go("/ch/settings", err="SECRET ni kiriting")
    try:
        await ch_track.ping()
        n = await ch_track.push_all()
        await ch_track.pull()
    except ch_track.TrackError as e:
        return go("/ch/settings", err=f"Saqlandi, lekin tekshiruv muvaffaqiyatsiz: {e}")
    return go("/ch/settings", msg=f"Worker ishlayapti. {n} ta kanal havolasi yuborildi")


@router.post("/ch/settings/worker-sync")
async def ch_worker_sync():
    try:
        await ch_track.push_all()
        n = await ch_track.pull()
    except ch_track.TrackError as e:
        return go("/ch/settings", err=str(e))
    return go("/ch/settings", msg=f"Bosishlar yangilandi ({n} qator)")
