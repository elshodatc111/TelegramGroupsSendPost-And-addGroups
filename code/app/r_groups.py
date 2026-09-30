"""Guruhlar: ro'yxat, teglar, qora ro'yxat, qoidalarni tekshirish, yangi guruh qidirish."""
import asyncio
import json
import uuid

from fastapi import APIRouter, Form, Request
from fastapi.responses import JSONResponse

from . import db
from .config import log
from .core import manager
from .limits import blacklist_add
from .r_posts import groups_context
from .web import acc_id, go, need_account, page

router = APIRouter()


def _ids(v):
    return [int(x) for x in v if str(x).lstrip("-").isdigit()]


@router.get("/groups")
async def groups_page(request: Request):
    aid = acc_id(request)
    if not aid:
        return page(request, "groups.html", groups=[], lists=[], tags={}, bl=[], st={}, prog=None, deltas={})
    groups, lists, tags = groups_context(aid)
    week = db.q("SELECT tg_id, MIN(day) d FROM group_members_log WHERE account_id=? AND day>=date('now','-7 day') GROUP BY tg_id", (aid,))
    deltas = {}
    for w in week:
        old = db.one("SELECT members FROM group_members_log WHERE account_id=? AND tg_id=? AND day=?", (aid, w["tg_id"], w["d"]))
        cur = next((g["members"] for g in groups if g["tg_id"] == w["tg_id"]), None)
        if old and old["members"] is not None and cur is not None:
            deltas[w["tg_id"]] = cur - old["members"]
    bl = db.q("SELECT * FROM blacklist WHERE account_id=? ORDER BY created_at DESC", (aid,))
    st = await manager.get(aid).status()
    return page(request, "groups.html", groups=groups, lists=lists, tags=tags, bl=bl, st=st,
                prog=manager.rules_progress.get(aid), deltas=deltas)


@router.post("/groups/sync")
async def groups_sync(request: Request):
    if (r := need_account(request)):
        return r
    try:
        n = await manager.get(acc_id(request)).fetch_groups()
    except Exception as e:
        log.exception("Guruhlarni yangilash")
        return go("/groups", err=f"Guruhlarni olib bo'lmadi: {e}")
    return go("/groups", msg=f"{n} ta guruh/kanal yangilandi")


@router.post("/groups/lists")
async def list_save(request: Request, name: str = Form(...), tg_ids: list[int] = Form(default=[])):
    aid = acc_id(request)
    name = name.strip()
    if not aid or not name or not tg_ids:
        return go("/groups", err="Ro'yxat nomini kiriting va kamida bitta guruh tanlang")
    ex = db.one("SELECT id FROM group_lists WHERE account_id=? AND name=?", (aid, name))
    lid = ex["id"] if ex else db.ex("INSERT INTO group_lists(account_id,name) VALUES(?,?)", (aid, name))
    db.ex("DELETE FROM group_list_items WHERE list_id=?", (lid,))
    db.many("INSERT OR IGNORE INTO group_list_items(list_id,tg_id) VALUES(?,?)", [(lid, i) for i in tg_ids])
    return go("/groups", msg=f"'{name}' ro'yxati saqlandi ({len(tg_ids)} ta guruh)")


@router.post("/groups/lists/{lid}/delete")
async def list_delete(lid: int):
    db.ex("DELETE FROM group_lists WHERE id=?", (lid,))
    return go("/groups", msg="Ro'yxat o'chirildi")


@router.post("/groups/tags")
async def tags_edit(request: Request, tg_ids: list[int] = Form(default=[]), tags: str = Form(""), mode: str = Form("add")):
    aid = acc_id(request)
    names = [t.strip().lower().lstrip("#") for t in tags.replace(";", ",").split(",") if t.strip()]
    if not aid or not tg_ids:
        return go("/groups", err="Guruhlarni tanlang")
    if mode == "clear":
        db.many("DELETE FROM group_tags WHERE account_id=? AND tg_id=?", [(aid, i) for i in tg_ids])
        return go("/groups", msg="Teglar tozalandi")
    if not names:
        return go("/groups", err="Teg nomini kiriting (masalan: toshkent, savdo)")
    if mode == "remove":
        db.many("DELETE FROM group_tags WHERE account_id=? AND tg_id=? AND tag=?", [(aid, i, t) for i in tg_ids for t in names])
        return go("/groups", msg="Teglar olib tashlandi")
    if mode == "set":
        db.many("DELETE FROM group_tags WHERE account_id=? AND tg_id=?", [(aid, i) for i in tg_ids])
    db.many("INSERT OR IGNORE INTO group_tags(account_id,tg_id,tag) VALUES(?,?,?)", [(aid, i, t) for i in tg_ids for t in names])
    return go("/groups", msg=f"{len(tg_ids)} ta guruhga teg qo'shildi")


@router.post("/groups/blacklist")
async def bl_add(request: Request, tg_ids: list[int] = Form(default=[]), reason: str = Form("Qo'lda qo'shildi")):
    aid = acc_id(request)
    if not aid or not tg_ids:
        return go("/groups", err="Guruhlarni tanlang")
    titles = {g["tg_id"]: g["title"] for g in db.q("SELECT tg_id,title FROM groups WHERE account_id=?", (aid,))}
    for i in tg_ids:
        blacklist_add(aid, i, titles.get(i, ""), reason or "Qo'lda qo'shildi")
    return go("/groups", msg=f"{len(tg_ids)} ta guruh qora ro'yxatga qo'shildi")


@router.post("/groups/blacklist/remove")
async def bl_remove(request: Request, tg_ids: list[int] = Form(default=[])):
    aid = acc_id(request)
    if aid and tg_ids:
        db.many("DELETE FROM blacklist WHERE account_id=? AND tg_id=?", [(aid, i) for i in tg_ids])
    return go("/groups", msg="Qora ro'yxatdan chiqarildi")


@router.post("/groups/ads-ok")
async def ads_ok(request: Request, tg_ids: list[int] = Form(default=[]), value: int = Form(1)):
    aid = acc_id(request)
    if aid and tg_ids:
        db.many("UPDATE groups SET ads_ok=? WHERE account_id=? AND tg_id=?", [(int(bool(value)), aid, i) for i in tg_ids])
    return go("/groups", msg="Saqlandi")


# ---------------- qoidalarni tekshirish (fon jarayoni) ----------------
async def _check_rules(aid: int):
    prog = manager.rules_progress[aid]
    svc = manager.get(aid)
    rows = db.q("SELECT tg_id FROM groups WHERE account_id=?", (aid,))
    prog.update(total=len(rows), done=0, fail=0, running=True)
    try:
        for r in rows:
            try:
                info = await svc.group_info(r["tg_id"])
                db.ex("UPDATE groups SET about=?, slowmode=?, no_media=?, no_links=?, ads_flag=?, checked_at=? "
                      "WHERE account_id=? AND tg_id=?", (info["about"], info["slowmode"], info["no_media"],
                                                         info["no_links"], info["ads_flag"], db.now(), aid, r["tg_id"]))
            except Exception as e:
                prog["fail"] += 1
                name = type(e).__name__
                if name == "FloodWaitError":
                    await asyncio.sleep(min(getattr(e, "seconds", 30) + 1, 120))
                log.info("Qoida tekshirish xatosi %s: %s", r["tg_id"], name)
            prog["done"] += 1
            await asyncio.sleep(1.2)
    finally:
        prog["running"] = False


@router.post("/groups/rules/start")
async def rules_start(request: Request):
    if (r := need_account(request)):
        return r
    aid = acc_id(request)
    if manager.rules_progress.get(aid, {}).get("running"):
        return go("/groups", err="Tekshiruv allaqachon ketmoqda")
    if not (await manager.get(aid).status()).get("authorized"):
        return go("/accounts", err="Akkaunt ulanmagan")
    manager.rules_progress[aid] = {"running": True, "total": 0, "done": 0, "fail": 0}
    asyncio.create_task(_check_rules(aid))
    return go("/groups", msg="Guruh qoidalari tekshirilmoqda (bir necha daqiqa ketishi mumkin)")


@router.get("/api/groups/rules-progress")
async def rules_progress(request: Request):
    return manager.rules_progress.get(acc_id(request)) or {"running": False, "total": 0, "done": 0, "fail": 0}


# ---------------- yangi guruhlarni topish ----------------
from .discovery import PRESET_SETS  # noqa: E402


def get_cfg_min(aid):
    from .auditor import get_cfg
    return get_cfg(aid)["min_members"]


@router.get("/discover")
async def discover(request: Request, q: str = ""):
    if (r := need_account(request)):
        return r
    aid = acc_id(request)
    results, error, hidden = [], None, {}
    if q.strip():
        try:
            from .auditor import get_cfg
            ac = get_cfg(aid)
            st = {}
            results = await manager.get(aid).search_public(q.strip(), strict=True, min_members=ac["min_members"],
                                                           check_ads=bool(ac["check_ads"]), stats=st)
            hidden = st.get("hidden", {})
        except Exception as e:
            error = f"Qidirishda xato: {e}"
    return page(request, "discover.html", q=q, results=results, error=error, hidden=hidden, preset_sets=PRESET_SETS, min_members=get_cfg_min(aid))


@router.post("/discover/preset-to-autojoin")
async def preset_to_autojoin(request: Request, set: str = Form("")):
    """Tanlangan tayyor so'zlar to'plamini Avto-topish qidiruv so'zlariga qo'shadi."""
    if (r := need_account(request)):
        return r
    from . import discovery
    if set not in discovery.PRESET_SETS:
        return go("/discover", err="Noma'lum to'plam")
    presets, category = discovery.PRESET_SETS[set]
    aid = acc_id(request)
    cfg = discovery.get_cfg(aid)
    kws = {k: list(v) for k, v in cfg["keywords"].items()}
    cur = kws.setdefault(category, [])
    have = {x.lower() for x in cur}
    add = []
    for ws in presets.values():
        for w in ws:
            if w.lower() not in have:
                have.add(w.lower())
                add.append(w)
    cur.extend(add)
    db.ex("UPDATE autojoin SET keywords_json=? WHERE account_id=?", (json.dumps(kws, ensure_ascii=False), aid))
    return go("/discover", msg=f"{len(add)} ta so'z Avto-topish qidiruviga qo'shildi ({category})")


@router.post("/discover/korean-to-autojoin")
async def korean_to_autojoin(request: Request):
    return await preset_to_autojoin(request, set="Koreys tili kurslari")


@router.post("/discover/join")
async def discover_join(request: Request, usernames: list[str] = Form(default=[]), q: str = Form("")):
    if (r := need_account(request)):
        return r
    aid = acc_id(request)
    items = [{"kind": "username", "key": u, "ref": "@" + u} for u in usernames if u]
    if not items:
        return go("/discover?q=" + q, err="Kamida bitta guruhni belgilang")
    cols = [{"label": f"Qidiruv: {q}", "items": items}]
    bid = db.ex("INSERT INTO join_batches(account_id,filename,status,created_at,raw) VALUES(?,?,?,?,?)",
                (aid, f"Qidiruv: {q}", "draft", db.now(), json.dumps(cols, ensure_ascii=False)))
    return go(f"/join/{bid}/setup")


# ---------------- Excel: eksport va import ----------------
def _xlsx_bytes(aid: int) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from .analytics import group_scores
    tags: dict[int, list[str]] = {}
    for r in db.q("SELECT tg_id, tag FROM group_tags WHERE account_id=? ORDER BY tag", (aid,)):
        tags.setdefault(r["tg_id"], []).append(r["tag"])
    lists: dict[int, list[str]] = {}
    for r in db.q("SELECT i.tg_id, l.name FROM group_list_items i JOIN group_lists l ON l.id=i.list_id WHERE l.account_id=?", (aid,)):
        lists.setdefault(r["tg_id"], []).append(r["name"])
    bl = {r["tg_id"]: r["reason"] for r in db.q("SELECT tg_id, reason FROM blacklist WHERE account_id=?", (aid,))}
    ok = {r["tg_id"]: r["ads_ok"] for r in db.q("SELECT tg_id, ads_ok FROM groups WHERE account_id=?", (aid,))}
    wb = Workbook()
    ws = wb.active
    ws.title = "Guruhlar"
    head = ["tg_id", "Nomi", "Username", "Turi", "A'zolar", "Teglar", "Ro'yxatlar", "Qora ro'yxat", "Qora ro'yxat sababi",
            "Reklama ruxsati", "Postlar", "O'rt. ko'rish", "Ball"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="DC2626")
        c.alignment = Alignment(vertical="center")
    for g in group_scores(aid, 365):
        ws.append([g["tg_id"], g["title"], g["username"] or "", g["kind"], g["members"], ", ".join(tags.get(g["tg_id"], [])),
                   ", ".join(lists.get(g["tg_id"], [])), "Ha" if g["tg_id"] in bl else "", bl.get(g["tg_id"], ""),
                   "Ha" if ok.get(g["tg_id"]) else "", g["sent"], g["avg_views"], g["score"]])
    for col, w in zip("ABCDEFGHIJKLM", [16, 36, 22, 12, 10, 26, 26, 12, 30, 14, 9, 12, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/groups/export.xlsx")
async def groups_export(request: Request):
    from fastapi.responses import Response
    if (r := need_account(request)):
        return r
    return Response(_xlsx_bytes(acc_id(request)),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="guruhlar.xlsx"'})


@router.post("/groups/import")
async def groups_import(request: Request):
    """Excel'dagi teglar, ro'yxatlar, qora ro'yxat va reklama ruxsatini mavjud guruhlarga qo'llaydi."""
    import io

    from openpyxl import load_workbook
    if (r := need_account(request)):
        return r
    aid = acc_id(request)
    form = await request.form()
    up = form.get("file")
    if not up or not getattr(up, "filename", "").lower().endswith(".xlsx"):
        return go("/groups", err="Faqat .xlsx fayl yuklang (guruhlar.xlsx ko'rinishida)")
    try:
        ws = load_workbook(io.BytesIO(await up.read()), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        return go("/groups", err=f"Faylni o'qib bo'lmadi: {e}")
    if not rows:
        return go("/groups", err="Fayl bo'sh")
    head = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            if n in head:
                return head.index(n)
        return None
    c_id, c_user = col("tg_id"), col("username")
    c_tags, c_lists = col("teglar"), col("ro'yxatlar")
    c_bl, c_reason, c_ads = col("qora ro'yxat"), col("qora ro'yxat sababi"), col("reklama ruxsati")
    if c_id is None and c_user is None:
        return go("/groups", err="Faylda 'tg_id' yoki 'Username' ustuni topilmadi")
    by_id = {g["tg_id"]: g for g in db.q("SELECT tg_id,title,username FROM groups WHERE account_id=?", (aid,))}
    by_user = {(g["username"] or "").lower(): g for g in by_id.values() if g["username"]}
    hit = miss = 0
    list_ids = {l["name"]: l["id"] for l in db.q("SELECT id,name FROM group_lists WHERE account_id=?", (aid,))}
    for row in rows[1:]:
        g = None
        if c_id is not None and row[c_id] not in (None, ""):
            try:
                g = by_id.get(int(row[c_id]))
            except (TypeError, ValueError):
                g = None
        if g is None and c_user is not None and row[c_user]:
            g = by_user.get(str(row[c_user]).strip().lstrip("@").lower())
        if g is None:
            miss += 1
            continue
        hit += 1
        tg = g["tg_id"]
        if c_tags is not None:
            db.ex("DELETE FROM group_tags WHERE account_id=? AND tg_id=?", (aid, tg))
            for t in str(row[c_tags] or "").replace(";", ",").split(","):
                if t.strip():
                    db.ex("INSERT OR IGNORE INTO group_tags(account_id,tg_id,tag) VALUES(?,?,?)", (aid, tg, t.strip().lower().lstrip("#")))
        if c_lists is not None:
            for name in [x.strip() for x in str(row[c_lists] or "").split(",") if x.strip()]:
                if name not in list_ids:
                    list_ids[name] = db.ex("INSERT INTO group_lists(account_id,name) VALUES(?,?)", (aid, name))
                db.ex("INSERT OR IGNORE INTO group_list_items(list_id,tg_id) VALUES(?,?)", (list_ids[name], tg))
        if c_bl is not None:
            if str(row[c_bl] or "").strip().lower() in ("ha", "yes", "1", "true"):
                blacklist_add(aid, tg, g["title"], str(row[c_reason] or "Excel'dan import"))
            else:
                db.ex("DELETE FROM blacklist WHERE account_id=? AND tg_id=?", (aid, tg))
        if c_ads is not None:
            db.ex("UPDATE groups SET ads_ok=? WHERE account_id=? AND tg_id=?",
                  (1 if str(row[c_ads] or "").strip().lower() in ("ha", "yes", "1", "true") else 0, aid, tg))
    return go("/groups", msg=f"Import tugadi: {hit} ta guruhga qo'llandi" + (f", {miss} tasi topilmadi (akkaunt a'zo emas)" if miss else ""))
