"""Statistika hisoboti: ma'lumot yig'ish, Excel va PDF eksport."""
import io
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import db

RED = "DC2626"


def stats_data(aid: int, days: int = 30) -> dict:
    since = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:%M:%S")
    base = "FROM job_targets jt JOIN jobs j ON j.id=jt.job_id WHERE j.account_id=? "
    sent = db.one(f"SELECT COUNT(*) c, COALESCE(SUM(jt.views),0) v, COALESCE(SUM(jt.reactions),0) r, "
                  f"COALESCE(SUM(jt.replies),0) p, COALESCE(SUM(jt.forwards),0) f {base} AND jt.status='sent' AND jt.sent_at>=?",
                  (aid, since))
    failed = db.one(f"SELECT COUNT(*) c {base} AND jt.status='failed' AND jt.sent_at>=?", (aid, since))["c"]
    deleted = db.one(f"SELECT COUNT(*) c {base} AND jt.deleted=1 AND jt.sent_at>=?", (aid, since))["c"]
    by_day = [dict(r) for r in db.q(
        f"SELECT substr(jt.sent_at,1,10) d, COUNT(*) sent, COALESCE(SUM(jt.views),0) views {base} AND jt.status='sent' "
        f"AND jt.sent_at>=? GROUP BY d ORDER BY d", (aid, since))]
    variants = [dict(r) for r in db.q(
        f"SELECT COALESCE(NULLIF(j.name,''),'Kampaniyasiz') name, COALESCE(jv.idx,0)+1 variant, COUNT(*) sent, "
        f"ROUND(AVG(jt.views),1) avg_views, COALESCE(SUM(jt.reactions),0) reactions, COALESCE(SUM(jt.replies),0) replies "
        f"FROM job_targets jt JOIN jobs j ON j.id=jt.job_id LEFT JOIN job_variants jv ON jv.id=jt.variant_id "
        f"WHERE j.account_id=? AND jt.status='sent' AND jt.sent_at>=? GROUP BY name, variant ORDER BY name, variant",
        (aid, since))]
    rows = db.q(
        f"SELECT jt.tg_id, MAX(jt.title) title, SUM(jt.status='sent') sent, SUM(jt.status='failed') failed, "
        f"SUM(jt.deleted) deleted, ROUND(AVG(CASE WHEN jt.status='sent' THEN jt.views END),1) avg_views, "
        f"COALESCE(SUM(jt.reactions),0) reactions, COALESCE(SUM(jt.replies),0) replies "
        f"{base} AND jt.sent_at>=? GROUP BY jt.tg_id", (aid, since))
    week = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d")
    groups = []
    for r in rows:
        d = dict(r)
        d["sent"], d["failed"], d["deleted"] = d["sent"] or 0, d["failed"] or 0, d["deleted"] or 0
        old = db.one("SELECT members FROM group_members_log WHERE account_id=? AND tg_id=? AND day>=? ORDER BY day LIMIT 1",
                     (aid, d["tg_id"], week))
        cur = db.one("SELECT members FROM group_members_log WHERE account_id=? AND tg_id=? ORDER BY day DESC LIMIT 1",
                     (aid, d["tg_id"]))
        d["members"] = cur["members"] if cur else None
        d["delta"] = (cur["members"] - old["members"]) if cur and old and cur["members"] is not None and old["members"] is not None else None
        if d["deleted"] >= 2 or d["failed"] >= 3:
            d["rating"], d["label"] = "bad", "Ishonchsiz"
        elif d["deleted"] or d["failed"]:
            d["rating"], d["label"] = "warn", "Ehtiyot"
        else:
            d["rating"], d["label"] = "good", "Yaxshi"
        d["bl"] = bool(db.one("SELECT 1 FROM blacklist WHERE account_id=? AND tg_id=?", (aid, d["tg_id"])))
        groups.append(d)
    order = {"bad": 0, "warn": 1, "good": 2}
    groups.sort(key=lambda g: (order[g["rating"]], -(g["avg_views"] or 0)))
    return {"days": days, "sent": sent["c"], "views": sent["v"], "reactions": sent["r"], "replies": sent["p"],
            "forwards": sent["f"], "failed": failed, "deleted": deleted, "by_day": by_day, "variants": variants,
            "groups": groups}


def export_xlsx(data: dict, account_name: str) -> bytes:
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=RED)

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for col, w in zip("ABCDEFGHIJ", widths):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Xulosa"
    sheet(ws, ["Ko'rsatkich", "Qiymat"], [
        ["Akkaunt", account_name], ["Davr (kun)", data["days"]], ["Yuborildi", data["sent"]], ["Xato", data["failed"]],
        ["Ko'rishlar", data["views"]], ["Reaksiyalar", data["reactions"]], ["Javoblar", data["replies"]],
        ["Ulashishlar", data["forwards"]], ["O'chirilgan postlar", data["deleted"]]], (28, 22))
    sheet(wb.create_sheet("Kunlik"), ["Sana", "Yuborildi", "Ko'rishlar"],
          [[d["d"], d["sent"], d["views"]] for d in data["by_day"]], (14, 12, 12))
    sheet(wb.create_sheet("Variantlar"), ["Kampaniya", "Variant", "Yuborildi", "O'rtacha ko'rish", "Reaksiya", "Javob"],
          [[v["name"], v["variant"], v["sent"], v["avg_views"], v["reactions"], v["replies"]] for v in data["variants"]],
          (28, 10, 12, 16, 12, 12))
    sheet(wb.create_sheet("Guruhlar reytingi"),
          ["Guruh", "Baho", "Yuborildi", "Xato", "O'chirilgan", "O'rtacha ko'rish", "Reaksiya", "Javob", "A'zolar", "7 kunlik o'zgarish"],
          [[g["title"], g["label"], g["sent"], g["failed"], g["deleted"], g["avg_views"], g["reactions"], g["replies"],
            g["members"], g["delta"]] for g in data["groups"]], (36, 14, 12, 8, 12, 16, 10, 8, 10, 16))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for p in (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\segoeui.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf"):
        if Path(p).exists():
            try:
                pdfmetrics.registerFont(TTFont("AppFont", p))
                return "AppFont"
            except Exception:
                continue
    return "Helvetica"


def export_pdf(data: dict, account_name: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName=font, textColor=colors.HexColor("#" + RED))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=font)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName=font)

    def table(rows, widths=None):
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + RED)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F8")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5D9E3")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        return t

    el = [Paragraph("Reklama hisoboti", h1),
          Paragraph(f"Akkaunt: {account_name} &nbsp;|&nbsp; Davr: oxirgi {data['days']} kun &nbsp;|&nbsp; {datetime.now():%Y-%m-%d %H:%M}", body),
          Spacer(1, 6 * mm),
          table([["Yuborildi", "Xato", "Ko'rishlar", "Reaksiyalar", "Javoblar", "O'chirilgan"],
                 [data["sent"], data["failed"], data["views"], data["reactions"], data["replies"], data["deleted"]]]),
          Spacer(1, 6 * mm)]
    if data["variants"]:
        el += [Paragraph("Variantlar samaradorligi (A/B)", h2),
               table([["Kampaniya", "Variant", "Yuborildi", "O'rt. ko'rish", "Reaksiya", "Javob"]] +
                     [[v["name"][:40], f"V{v['variant']}", v["sent"], v["avg_views"] or 0, v["reactions"], v["replies"]] for v in data["variants"]]),
               Spacer(1, 6 * mm)]
    el.append(Paragraph("Guruhlar reytingi", h2))
    rows = [["Guruh", "Baho", "Yuborildi", "Xato", "O'chirilgan", "O'rt. ko'rish", "Reaksiya", "Javob", "A'zolar"]]
    for g in data["groups"][:80]:
        rows.append([(g["title"] or "")[:44], g["label"], g["sent"], g["failed"], g["deleted"], g["avg_views"] or 0,
                     g["reactions"], g["replies"], g["members"] or ""])
    el.append(table(rows, [80 * mm, 26 * mm, 22 * mm, 16 * mm, 22 * mm, 24 * mm, 20 * mm, 18 * mm, 22 * mm]))
    doc.build(el)
    return buf.getvalue()
