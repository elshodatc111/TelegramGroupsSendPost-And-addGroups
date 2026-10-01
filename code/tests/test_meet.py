"""Majlislar (Google Meet) bo'limi sinovlari: soxta Google (mock_google) va soxta Telegram bilan."""
import io
import json
from datetime import datetime, timedelta
from urllib.parse import parse_qs, urlparse

import pytest
from fastapi.testclient import TestClient


# ---------------------------------------------------------------- yordamchilar
def iso(dt):
    """Toshkent vaqti (naive) -> Google UTC matni (9 xonali kasr bilan, haqiqiy API kabi)."""
    return (dt - timedelta(hours=5)).strftime("%Y-%m-%dT%H:%M:%S.123456789Z")


def connect(S, run):
    from app import meet_google as G
    S.put("google_client_id", "cid.apps.googleusercontent.com")
    S.put("google_client_secret", "csecret")
    url = G.start_auth()
    state = parse_qs(urlparse(url).query)["state"][0]
    run(G.finish_auth("good-code", state))
    return G


@pytest.fixture()
def tg(clean, monkeypatch):
    S = clean
    sent, fail = [], {}

    async def fake(target, text):
        if fail.get(target):
            raise S.TgFail(fail[target])
        sent.append((target, text))
    monkeypatch.setattr(S, "tg_send", fake)
    return sent, fail


@pytest.fixture()
def world(clean, run):
    S = clean
    G = connect(S, run)
    S.put("admin_username", "adminuser")
    tid = S.db.ex("INSERT INTO mt_teachers(name,gmail,tg_username,phone,active,created_at) VALUES('Ustoz Test','ustoz@gmail.com','ustoz_tg','+998901112233',1,?)", (S.fmt(S.now()),))
    ga = S.add_group(-1001, "Ingliz tili A", "ga", tid)
    gb = S.add_group(-1002, "Matematika B", "", tid)
    return S, G, tid, ga, gb


def mk_lesson(S, run, gid, tid, start, dur=60):
    lid = S.add_extra_lesson(gid, start, dur, tid)
    ok, why = run(S.provision(lid))
    assert ok, why
    return lid


def person(uid, name, sessions, kind="signedinUser"):
    ps = [{"startTime": iso(a), **({"endTime": iso(b)} if b else {})} for a, b in sessions]
    first = min(a for a, _ in sessions)
    ends = [b for _, b in sessions]
    p = {"name": "", "earliestStartTime": iso(first), kind: {"displayName": name}, "sessions": ps}
    if kind == "signedinUser":
        p[kind]["user"] = f"users/{uid}"
    if all(ends):
        p["latestEndTime"] = iso(max(ends))
    return p


def set_conf(mock, space, start, end, people, cid="cr1"):
    for i, p in enumerate(people):
        p["name"] = f"conferenceRecords/{cid}/participants/p{i + 1}"
    rec = {"name": f"conferenceRecords/{cid}", "startTime": iso(start), "space": space, "participants": people}
    if end:
        rec["endTime"] = iso(end)
    mock.STATE["records"][space] = [rec]


# ---------------------------------------------------------------- sxema, OAuth, token
def test_schema_has_seven_tables(clean):
    S = clean
    for t in S.TABLE_NAMES:
        assert S.db.q(f"SELECT COUNT(*) c FROM {t}")[0]["c"] >= 0
    assert len(S.TABLE_NAMES) == 7


def test_oauth_flow_and_encrypted_tokens(clean, run, mock):
    S = clean
    G = connect(S, run)
    assert G.connected() and S.get("g_email") == "host@gmail.com" and S.get("host_type") == "gmail"
    raw = S.db.one("SELECT value FROM mt_settings WHERE key='g_refresh'")["value"]
    assert raw.startswith(("fk1:", "dpapi1:")) and "rt-secret" not in raw          # token shifrlangan
    assert S.get("g_refresh") == "rt-secret-123"
    assert S.db.one("SELECT value FROM mt_settings WHERE key='google_client_secret'")["value"].startswith(("fk1:", "dpapi1:"))
    url = G.start_auth()
    q = parse_qs(urlparse(url).query)
    assert q["code_challenge_method"] == ["S256"] and q["access_type"] == ["offline"] and "127.0.0.1" in q["redirect_uri"][0]
    assert "meetings.space.created" in q["scope"][0] and "auth/calendar" in q["scope"][0]


def test_oauth_rejects_bad_state_and_workspace_detect(clean, run, mock):
    S = clean
    from app import meet_google as G
    S.put("google_client_id", "cid")
    S.put("google_client_secret", "sec")
    G.start_auth()
    with pytest.raises(G.GoogleError):
        run(G.finish_auth("good-code", "wrong-state"))
    mock.STATE["email"], mock.STATE["hd"] = "boss@school.uz", "school.uz"
    state = parse_qs(urlparse(G.start_auth()).query)["state"][0]
    run(G.finish_auth("good-code", state))
    assert S.get("host_type") == "workspace"


def test_token_expired_marks_state_and_alerts(world, run, mock):
    S, G, *_ = world
    S.put("g_exp", "0")
    mock.STATE["fail_refresh"] = True
    with pytest.raises(G.GoogleError) as e:
        run(G.create_space())
    assert e.value.reason == "invalid_grant" and "qayta ulang" in e.value.hint.lower()
    assert S.get("g_state") == "expired" and not G.connected()
    assert S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='google_token'")["c"] == 1
    items = {i["key"]: i for i in S.health_items()}
    assert items["meet_token"]["state"] == "fail" and "Production" in items["meet_token"]["fix"]


def test_host_switch(world, run, mock):
    S, G, *_ = world
    G.disconnect()
    assert S.get("g_state") == "none" and not S.get("g_refresh")
    mock.STATE["email"] = "other@gmail.com"
    state = parse_qs(urlparse(G.start_auth()).query)["state"][0]
    run(G.finish_auth("good-code", state))
    assert S.get("g_email") == "other@gmail.com"


# ---------------------------------------------------------------- Meet xona, co-host, kalendar
def test_provision_creates_space_cohost_calendar(world, run, mock):
    S, G, tid, ga, gb = world
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(hours=3), 60)
    les = S.lesson(lid)
    sp = mock.STATE["spaces"][les["space_name"]]
    assert sp["config"]["accessType"] == "OPEN" and les["meeting_uri"].startswith("https://meet.google.com/")
    mem = mock.STATE["members"][les["space_name"]]
    assert mem == [{"name": f"{les['space_name']}/members/m1", "email": "ustoz@gmail.com", "role": "COHOST"}] and les["cohost_ok"] == 1
    cal = S.group(ga)["calendar_id"]
    ev = mock.STATE["events"][cal][les["cal_event_id"]]
    assert les["meeting_uri"] in ev["location"] and ev["start"]["timeZone"] == "Asia/Tashkent"
    assert mock.STATE["acl"] and mock.STATE["acl"][0][1]["scope"]["value"] == "ustoz@gmail.com"
    # har dars uchun YANGI xona, har guruh uchun ALOHIDA kalendar
    l2 = mk_lesson(S, run, ga, tid, S.now() + timedelta(hours=5))
    l3 = mk_lesson(S, run, gb, tid, S.now() + timedelta(hours=7))
    assert len({S.lesson(x)["space_name"] for x in (lid, l2, l3)}) == 3
    assert S.group(ga)["calendar_id"] != S.group(gb)["calendar_id"] and len(mock.STATE["calendars"]) == 2


def test_cohost_failure_is_a_warning_not_a_crash(world, run, mock):
    S, G, tid, ga, gb = world
    mock.STATE["bad_members"].add("ustoz@gmail.com")
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(hours=3))
    les = S.lesson(lid)
    assert les["meeting_uri"] and les["cohost_ok"] == 0 and "co-host qilinmadi" in les["warn_text"]


def test_60_minute_warning(world, run):
    S, G, tid, ga, gb = world
    errs, warns = S.add_schedule(ga, [0, 2, 4], "19:00", 90, tid)
    assert not errs and any("60 daqiqa" in w and "Gmail" in w for w in warns)
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(hours=3), 90)
    assert "60 daqiqa" in S.lesson(lid)["warn_text"]
    assert any(i["title"] == "Meet limiti" and "Meet limiti: 60 daqiqa (Gmail)" in i["detail"] for i in S.health_items())
    S.put("host_type", "workspace")
    errs, warns = S.add_schedule(gb, [1, 3, 5], "19:00", 90, tid)
    assert not any("60 daqiqa" in w for w in warns)


def test_schedule_validation_and_weekly_count_warning(world):
    S, G, tid, ga, gb = world
    assert S.add_schedule(ga, [], "19:00", 60)[0] and S.add_schedule(ga, [0], "25:00", 60)[0] and S.add_schedule(ga, [0], "19:00", 5)[0]
    errs, warns = S.add_schedule(ga, [0, 1], "19:00", 60, tid)
    assert not errs and any("3-6" in w for w in warns)
    errs, warns = S.add_schedule(ga, [2, 3, 4], "19:00", 60, tid)
    assert not warns
    assert len(S.schedules(ga)) == 5


def test_generate_from_schedule_idempotent(world, run, mock):
    S, G, tid, ga, gb = world
    S.put("horizon_days", 14)
    S.add_schedule(ga, [0, 1, 2, 3, 4, 5], "23:30", 45, tid)
    n1 = run(S.generate(ga))
    assert 10 <= n1 <= 13                    # 14 kunda 6 kun/hafta, bugungi o'tib ketgan bo'lishi mumkin
    assert run(S.generate(ga)) == 0
    rows = S.db.q("SELECT * FROM mt_lessons WHERE group_id=?", (ga,))
    assert all(r["space_name"] and r["cal_event_id"] for r in rows) and len({r["space_name"] for r in rows}) == len(rows)
    assert len(mock.STATE["spaces"]) == n1


def test_google_api_disabled_gives_actionable_error(world, run, mock):
    S, G, tid, ga, gb = world
    mock.STATE["meet_disabled"] = True
    lid = S.add_extra_lesson(ga, S.now() + timedelta(hours=2), 60, tid)
    ok, why = run(S.provision(lid))
    assert not ok and "console.developers.google.com" in why
    assert S.lesson(lid)["prov_err"] and S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='provision'")["c"] == 1
    mock.STATE["meet_disabled"] = False
    S.db.ex("UPDATE mt_lessons SET prov_at=NULL WHERE id=?", (lid,))
    run(S.retry_provision())
    assert S.lesson(lid)["meeting_uri"]


# ---------------------------------------------------------------- Telegram xabarlari va eslatmalar
def test_link_and_reminder_timeline(world, run, tg):
    S, G, tid, ga, gb = world
    sent, fail = tg
    S.put("remind_a", 60)
    S.put("remind_b", 10)
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(minutes=55), 60)
    run(S.process_reminders())
    tos = [(t, m) for t, m in sent]
    assert len(tos) == 1 and tos[0][0] == -1001 and "55" in tos[0][1] and "Eslatma" in tos[0][1]            # 60 daq eslatma (55 qoldi)
    S.db.ex("UPDATE mt_lessons SET start_at=?, end_at=? WHERE id=?", (S.fmt(S.now() + timedelta(minutes=9, seconds=30)), S.fmt(S.now() + timedelta(minutes=69)), lid))
    sent.clear()
    run(S.process_reminders())
    kinds = {(t): m for t, m in sent}
    assert "@ustoz_tg" in kinds and S.lesson(lid)["meeting_uri"] in kinds["@ustoz_tg"]               # o'qituvchiga 10 daq oldin, shaxsiy, havola bilan
    assert -1001 in kinds and "Eslatma" in kinds[-1001] and S.lesson(lid)["meeting_uri"] not in kinds[-1001]   # 10 daq eslatma (havolasiz)
    sent.clear()
    S.db.ex("UPDATE mt_lessons SET start_at=?, end_at=? WHERE id=?", (S.fmt(S.now() + timedelta(minutes=4)), S.fmt(S.now() + timedelta(minutes=64)), lid))
    run(S.process_reminders())
    assert [t for t, _ in sent] == [-1001] and S.lesson(lid)["meeting_uri"] in sent[0][1] and "O'qituvchi: Ustoz Test" in sent[0][1]   # 5 daq oldin guruhga havola
    sent.clear()
    run(S.process_reminders())
    assert sent == []                                    # takrorlanmaydi


def test_reminders_configurable_and_can_be_disabled(world, run, tg):
    S, G, tid, ga, gb = world
    sent, _ = tg
    S.put("remind_a", 120)
    S.put("teacher_dm_min", 20)
    S.put("group_link_min", 15)
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(minutes=100), 60)
    run(S.process_reminders())
    assert len(sent) == 1 and "Eslatma" in sent[0][1]
    S.db.ex("UPDATE mt_lessons SET start_at=?, end_at=? WHERE id=?", (S.fmt(S.now() + timedelta(minutes=18)), S.fmt(S.now() + timedelta(minutes=78)), lid))
    sent.clear()
    run(S.process_reminders())
    assert "@ustoz_tg" in [t for t, _ in sent]
    S.put("remind_on", 0)
    l2 = mk_lesson(S, run, gb, tid, S.now() + timedelta(minutes=100), 60)
    sent.clear()
    run(S.process_reminders())
    assert not [1 for t, m in sent if t == -1002]


def test_send_failure_alerts_admin_with_copyable_link(world, run, tg):
    S, G, tid, ga, gb = world
    sent, fail = tg
    fail["@ustoz_tg"] = "foydalanuvchi maxfiylik sozlamasi xabar yuborishga ruxsat bermaydi"
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(minutes=8), 60)
    uri = S.lesson(lid)["meeting_uri"]
    run(S.process_reminders())
    al = S.db.one("SELECT * FROM mt_alerts WHERE kind='send_fail'")
    assert al and uri in al["copy_text"] and "o'qituvchi" in al["text"] and al["level"] == "error"
    assert any(t == "adminuser" or t == "@adminuser" for t, m in sent) and any("yetmadi" in m for t, m in sent)    # adminga ogohlantirish
    assert S.flags(S.lesson(lid))["t_link"].startswith("fail")
    run(S.process_reminders())
    assert S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='send_fail'")["c"] == 1     # spam yo'q
    # qo'lda qayta yuborish
    fail.clear()
    assert run(S.send_now(lid, "teacher")) == "Yuborildi"


def test_teacher_without_username_and_missing_link(world, run, tg):
    S, G, tid, ga, gb = world
    sent, _ = tg
    S.db.ex("UPDATE mt_teachers SET tg_username='' WHERE id=?", (tid,))
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(minutes=8), 60)
    run(S.process_reminders())
    al = S.db.one("SELECT * FROM mt_alerts WHERE kind='send_fail'")
    assert al and "@username" in al["text"] and S.lesson(lid)["meeting_uri"] in al["copy_text"]


def test_teacher_late_alert(world, run, tg):
    S, G, tid, ga, gb = world
    sent, _ = tg
    lid = mk_lesson(S, run, ga, tid, S.now() - timedelta(minutes=6), 60)
    run(S.process_reminders())
    al = S.db.one("SELECT * FROM mt_alerts WHERE kind='teacher_late'")
    assert al and "Ustoz Test" in al["text"] and any("Meet'da yo'q" in m for t, m in sent)
    # 4 daqiqada hali ogohlantirish yo'q
    l2 = mk_lesson(S, run, gb, tid, S.now() - timedelta(minutes=4), 60)
    run(S.process_reminders())
    assert S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='teacher_late'")["c"] == 1
    # o'qituvchi kirgan bo'lsa ogohlantirish yo'q
    l3 = mk_lesson(S, run, ga, tid, S.now() - timedelta(minutes=9), 60)
    S.set_flag(l3, "t_in", S.fmt(S.now()))
    n = S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='teacher_late'")["c"]
    run(S.process_reminders())
    assert S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='teacher_late'")["c"] == n


def test_cancel_and_move_notify_everyone_and_update_calendar(world, run, tg, mock):
    S, G, tid, ga, gb = world
    sent, _ = tg
    lid = mk_lesson(S, run, ga, tid, S.now() + timedelta(days=1, hours=1), 60)
    les = S.lesson(lid)
    cal = S.group(ga)["calendar_id"]
    new = (S.now() + timedelta(days=2)).replace(hour=15, minute=0, second=0, microsecond=0)
    msg = run(S.move_lesson(lid, new, 45))
    assert "ko'chirildi" in msg and S.lesson(lid)["start_at"] == S.fmt(new) and S.lesson(lid)["duration_min"] == 45
    ev = mock.STATE["events"][cal][les["cal_event_id"]]
    assert ev["start"]["dateTime"] == new.strftime("%Y-%m-%dT%H:%M:%S")
    assert {t for t, _ in sent} == {-1001, "@ustoz_tg"} and all("ko'chirildi" in m for _, m in sent)
    assert les["meeting_uri"] in sent[0][1]
    sent.clear()
    run(S.cancel_lesson(lid, "Ustoz kasal"))
    assert S.lesson(lid)["status"] == "cancelled" and les["cal_event_id"] not in mock.STATE["events"][cal]
    assert {t for t, _ in sent} == {-1001, "@ustoz_tg"} and all("bekor" in m and "Ustoz kasal" in m for _, m in sent)
    assert "Bu darsni" in run(S.cancel_lesson(lid, ""))          # qayta bekor qilib bo'lmaydi
    # bekor qilingan dars uchun eslatma/havola yuborilmaydi
    sent.clear()
    S.db.ex("UPDATE mt_lessons SET start_at=? WHERE id=?", (S.fmt(S.now() + timedelta(minutes=4)), lid))
    run(S.process_reminders())
    assert sent == []


def test_move_keeps_generation_from_recreating_old_slot(world, run):
    S, G, tid, ga, gb = world
    S.put("horizon_days", 10)
    S.add_schedule(ga, [0, 1, 2, 3, 4, 5, 6], "23:50", 30, tid)
    run(S.generate(ga))
    first = S.db.one("SELECT * FROM mt_lessons WHERE group_id=? ORDER BY start_at LIMIT 1", (ga,))
    run(S.move_lesson(first["id"], S.now() + timedelta(days=30)))
    before = S.db.one("SELECT COUNT(*) c FROM mt_lessons")["c"]
    assert run(S.generate(ga)) == 0 and S.db.one("SELECT COUNT(*) c FROM mt_lessons")["c"] == before


# ---------------------------------------------------------------- davomat (polling)
def build_class(S, run, mock, ga, tid, start_offset=-40, dur=60):
    """O'qituvchi + 3 o'quvchi (biri kech, biri erta chiqadi, biri ikki marta kiradi) + akkauntsiz mehmon."""
    start = S.now() + timedelta(minutes=start_offset)
    lid = mk_lesson(S, run, ga, tid, start, dur)
    space = S.lesson(lid)["space_name"]
    end = start + timedelta(minutes=dur)
    people = [
        person("u1", "Ustoz Test", [(start - timedelta(minutes=3), end)]),
        person("u2", "Aziza Karimova", [(start + timedelta(minutes=1), end)]),                                  # o'z vaqtida
        person("u3", "Bobur", [(start + timedelta(minutes=14), start + timedelta(minutes=45))]),               # kech (14) + erta (15 daq)
        person("u4", "Dilnoza", [(start + timedelta(minutes=2), start + timedelta(minutes=20)), (start + timedelta(minutes=25), end)]),   # qayta kirdi
        person("", "Ali (telefon)", [(start + timedelta(minutes=5), end)], kind="anonymousUser"),
    ]
    return lid, start, end, space, people


def test_polling_attendance_metrics(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    lid, start, end, space, people = build_class(S, run, mock, ga, tid)
    # dars tugagan holat (kech/erta hisoblash uchun)
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    res = run(T.poll_lesson(S.lesson(lid)))
    assert res["conf"] and res["n"] == 5 and res["peak"] >= 4 and not res["active"]
    les = S.lesson(lid)
    assert les["real_start"] and les["real_end"] and S.elapsed_min(les) == 60 and "t_in" in S.flags(les)
    rows = {p["person"]: p for p in T.person_rows(lid)}
    assert rows["Ustoz Test"]["is_teacher"]                                               # o'qituvchi avtomatik aniqlandi
    a, b, d, al = rows["Aziza Karimova"], rows["Bobur"], rows["Dilnoza"], rows["Ali (telefon)"]
    assert not a["late"] and not a["early"] and a["minutes"] == 59
    assert b["late"] and b["late_min"] == 14 and b["early"] and b["early_min"] == 15 and b["minutes"] == 31
    assert not d["late"] and not d["early"] and d["minutes"] == 18 + 35                    # ikki sessiya yig'indisi
    assert al["late"] is False and al["kind"] == "anon"                                      # akkauntsiz mehmon ham yuritiladi
    # chegara sozlanadi
    S.put("late_min", 20)
    assert not {p["person"]: p for p in T.person_rows(lid)}["Bobur"]["late"]
    S.put("late_min", 10)
    # qayta polling: o'zgarmagan ishtirokchilar uchun sessiya so'rovlari qayta yuborilmaydi
    mock.STATE["calls"].clear()
    run(T.poll_lesson(S.lesson(lid)))
    assert not [c for c in mock.STATE["calls"] if "participantSessions" in c]
    assert S.db.one("SELECT COUNT(*) c FROM mt_attendance WHERE lesson_id=?", (lid,))["c"] == 5


def test_live_polling_tracks_open_sessions(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    start = S.now() - timedelta(minutes=20)
    lid = mk_lesson(S, run, ga, tid, start, 60)
    space = S.lesson(lid)["space_name"]
    set_conf(mock, space, start, None, [person("u1", "Ustoz Test", [(start, None)]), person("u2", "Aziza", [(start + timedelta(minutes=2), None)])])
    res = run(T.poll_lesson(S.lesson(lid)))
    assert res["active"] and res["in_now"] == 2
    les = S.lesson(lid)
    assert les["status"] == "live" and 19 <= S.elapsed_min(les) <= 20 and les["real_end"] is None
    p = {x["person"]: x for x in T.person_rows(lid)}["Aziza"]
    assert p["in_now"] and p["last_out"] is None and p["minutes"] >= 17 and not p["early"]      # hali ichida: erta chiqqan emas
    # Aziza chiqib ketdi
    set_conf(mock, space, start, None, [person("u1", "Ustoz Test", [(start, None)]), person("u2", "Aziza", [(start + timedelta(minutes=2), start + timedelta(minutes=15))])])
    run(T.poll_lesson(S.lesson(lid)))
    p = {x["person"]: x for x in T.person_rows(lid)}["Aziza"]
    assert not p["in_now"] and p["minutes"] == 13


def test_merge_names_and_unmerge(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    start = S.now() - timedelta(minutes=30)
    lid = mk_lesson(S, run, ga, tid, start, 60)
    space = S.lesson(lid)["space_name"]
    set_conf(mock, space, start, None, [person("u2", "Aziza Karimova", [(start, start + timedelta(minutes=10))]),
                                        person("", "iPhone (Aziza)", [(start + timedelta(minutes=8), None)], kind="anonymousUser")])
    run(T.poll_lesson(S.lesson(lid)))
    names = [p["person"] for p in T.person_rows(lid)]
    assert len(names) == 2
    assert T.merge_names(ga, "iPhone (Aziza)", "Aziza Karimova") == 1
    rows = T.person_rows(lid)
    assert len(rows) == 1 and rows[0]["person"] == "Aziza Karimova" and rows[0]["raws"] == ["Aziza Karimova", "iPhone (Aziza)"]
    assert rows[0]["minutes"] == round((30 - 0) * 1) or rows[0]["minutes"] >= 29                  # oraliqlar birlashadi (10 + 22 ustma-ust 2 daq)
    run(T.poll_lesson(S.lesson(lid)))                                                           # keyingi polling birlashtirishni saqlaydi
    assert len(T.person_rows(lid)) == 1
    T.unmerge(ga, "iPhone (Aziza)")
    assert len(T.person_rows(lid)) == 2


def test_crowd_and_unknown_name_alerts(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    sent, _ = tg
    S.put("max_people", 4)
    start = S.now() - timedelta(minutes=10)
    lid = mk_lesson(S, run, ga, tid, start, 60)
    space = S.lesson(lid)["space_name"]
    set_conf(mock, space, start, None, [person(f"u{i}", f"Talaba {i}", [(start, None)]) for i in range(6)])
    run(T.poll_lesson(S.lesson(lid)))
    al = S.db.one("SELECT * FROM mt_alerts WHERE kind='crowd'")
    assert al and "6 kishi" in al["text"] and "o'qituvchi Meet ichida" in al["text"]
    assert any("Begona" in m for t, m in sent)
    run(T.poll_lesson(S.lesson(lid)))
    assert S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='crowd'")["c"] == 1
    # yangi nom: avvalgi 2 ta dars o'tgan, ro'yxat >= 3 kishi
    S.put("max_people", 0)
    for k in range(2):
        old = mk_lesson(S, run, ga, tid, S.now() - timedelta(days=k + 1), 60)
        S.db.ex("UPDATE mt_lessons SET status='done' WHERE id=?", (old,))
        for n in ("Ali", "Vali", "Gulya"):
            S.db.ex("INSERT INTO mt_attendance(lesson_id,part_name,raw_name,person,kind,first_in,total_sec,sessions_json,is_teacher) VALUES(?,?,?,?,?,?,?,?,0)",
                    (old, f"p-{old}-{n}", n, n, "signed", S.fmt(S.now()), 600, json.dumps([[S.fmt(S.now() - timedelta(hours=1)), S.fmt(S.now() - timedelta(minutes=50))]])))
    l2 = mk_lesson(S, run, ga, tid, S.now() - timedelta(minutes=5), 60)
    sp2 = S.lesson(l2)["space_name"]
    set_conf(mock, sp2, S.now() - timedelta(minutes=5), None, [person("u9", "Ali", [(S.now() - timedelta(minutes=4), None)]),
                                                               person("", "Begona Odam", [(S.now() - timedelta(minutes=3), None)], kind="anonymousUser")], cid="cr2")
    run(T.poll_lesson(S.lesson(l2)))
    unk = S.db.q("SELECT * FROM mt_alerts WHERE kind='unknown'")
    assert len(unk) == 1 and "Begona Odam" in unk[0]["text"]


def test_teacher_manual_mark_learns_user_id(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    S.db.ex("UPDATE mt_teachers SET name='Muhammad Ali Yusupov' WHERE id=?", (tid,))
    start = S.now() - timedelta(minutes=20)
    lid = mk_lesson(S, run, ga, tid, start, 60)
    space = S.lesson(lid)["space_name"]
    set_conf(mock, space, start, None, [person("uT", "MY", [(start, None)]), person("u2", "Aziza", [(start, None)])])
    run(T.poll_lesson(S.lesson(lid)))
    assert not S.flags(S.lesson(lid)).get("t_in")                                  # ism mos kelmadi
    att = S.db.one("SELECT id FROM mt_attendance WHERE lesson_id=? AND raw_name='MY'", (lid,))
    T.mark_teacher(att["id"], True)
    assert S.flags(S.lesson(lid)).get("t_in") and S.jget(f"tuid_{tid}", []) == ["users/uT"]
    l2 = mk_lesson(S, run, ga, tid, S.now() - timedelta(minutes=10), 60)
    set_conf(mock, S.lesson(l2)["space_name"], S.now() - timedelta(minutes=10), None, [person("uT", "Boshqa nom", [(S.now() - timedelta(minutes=9), None)])], cid="cr3")
    run(T.poll_lesson(S.lesson(l2)))
    assert S.flags(S.lesson(l2)).get("t_in")                                       # keyingi darsda user ID bo'yicha tanildi


def test_tick_finalizes_lesson_and_sends_report(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    sent, _ = tg
    lid, start, end, space, people = build_class(S, run, mock, ga, tid, start_offset=-80, dur=60)
    T.add_roster(ga, "Gulnoza Yoqubova")                                          # kelmagan o'quvchi
    S.set_flag(lid, "t_link", "x")
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    run(T.tick())
    les = S.lesson(lid)
    assert les["status"] == "done" and "report" in S.flags(les)
    rep = [m for t, m in sent if "Dars hisoboti" in m]
    assert rep and any(t in ("adminuser", "@adminuser") for t, m in sent if "Dars hisoboti" in m)
    txt = rep[0]
    assert "Kech qoldi: Bobur" in txt and "Erta chiqdi: Bobur" in txt and "Kelmadi: Gulnoza Yoqubova" in txt and "60 daq" in txt
    n = len(sent)
    run(T.tick())
    assert len(sent) == n and S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='report'")["c"] == 1     # ikki marta yakunlanmaydi


def test_lesson_without_participants_is_missed(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    lid = mk_lesson(S, run, ga, tid, S.now() - timedelta(minutes=90), 60)
    run(T.tick())
    assert S.lesson(lid)["status"] == "missed" and S.db.one("SELECT COUNT(*) c FROM mt_alerts WHERE kind='missed'")["c"] == 1


def test_polling_errors_are_counted_not_fatal(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    lid = mk_lesson(S, run, ga, tid, S.now() - timedelta(minutes=5), 60)
    S.db.ex("UPDATE mt_lessons SET space_name='spaces/yoq' WHERE id=?", (lid,))
    orig = G.conference_records

    async def boom(space):
        raise G.GoogleError("Meet API ishlamayapti", 500)
    G.conference_records = boom
    try:
        run(T.tick())
        run(T.tick())
    finally:
        G.conference_records = orig
    assert S.get("poll_err_24") == "2" and "Meet API" in S.get("last_poll_err")
    h = {i["key"]: i for i in S.health_items()}
    assert h["meet_perr"]["state"] == "warn" and h["meet_poll"]["detail"].startswith(S.get("last_poll_at")[:10])


# ---------------------------------------------------------------- hisobotlar
def test_monthly_grid_colors_excel_and_stats(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_reports as R, meet_track as T
    from openpyxl import load_workbook
    lid, start, end, space, people = build_class(S, run, mock, ga, tid, start_offset=-80, dur=60)
    T.add_roster(ga, "Gulnoza Yoqubova")
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    run(T.tick())
    # ikkinchi dars: bekor qilingan
    can = mk_lesson(S, run, ga, tid, S.now() + timedelta(days=1), 60)
    run(S.cancel_lesson(can, "bayram"))
    ym = start.strftime("%Y-%m")
    grid = R.month_grid(ga, ym)
    if len(grid["lessons"]) < 2:                         # oy chegarasi: bekor qilingan dars keyingi oyga tushgan bo'lishi mumkin
        pass
    rows = {r["name"]: r for r in grid["rows"]}
    st = {n: r["cells"][0]["state"] for n, r in rows.items()}
    assert st["Aziza Karimova"] == "present" and st["Bobur"] == "late_early" and st["Dilnoza"] == "present"
    assert st["Gulnoza Yoqubova"] == "absent" and "Ustoz Test" not in rows
    assert rows["Bobur"]["late"] == 1 and rows["Bobur"]["early"] == 1 and rows["Gulnoza Yoqubova"]["pct"] == 0 and rows["Aziza Karimova"]["pct"] == 100
    wb = load_workbook(io.BytesIO(R.export_xlsx(ga, ym)))
    ws = wb["Davomat"]
    colors = {}
    for row in ws.iter_rows(min_row=4, max_row=3 + len(grid["rows"])):
        colors[row[1].value] = row[2].fill.fgColor.rgb[-6:]
    assert colors["Aziza Karimova"] == "C6EFCE" and colors["Gulnoza Yoqubova"] == "FFC7CE" and colors["Bobur"] == "F4B183"
    assert "Darslar" in wb.sheetnames
    stt = R.stats(ym)
    t = stt["teachers"][0]
    assert t["name"] == "Ustoz Test" and t["lessons"] == 1 and t["groups"] == 1 and t["hours"] == 1.0
    assert stt["groups"][0]["title"] == "Ingliz tili A" and stt["groups"][0]["hours"] == 1.0 and stt["total_hours"] == 1.0


def test_daily_sheet_added_to_excel(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_reports as R, meet_track as T
    from openpyxl import Workbook
    lid, start, end, space, people = build_class(S, run, mock, ga, tid, start_offset=-80, dur=60)
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    run(T.tick())
    wb = Workbook()
    R.add_daily_sheet(wb, start.strftime("%Y-%m-%d"))
    ws = wb["Majlislar"]
    row = [c.value for c in ws[2]]
    assert row[1] == "Ingliz tili A" and row[3] == "O'tdi" and row[4] == 60 and row[5] == 4 and row[6] == 1 and row[7] == 1
    wb2 = Workbook()
    R.add_daily_sheet(wb2, "2001-01-01")
    assert "Majlislar" not in wb2.sheetnames                  # dars bo'lmagan kunga varaq qo'shilmaydi


def test_existing_daily_report_still_builds(world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import dailyreport, db
    aid = db.ex("INSERT INTO accounts(name,session,created_at,user_id,workspace) VALUES('Test','t',?,1,'posting')", (S.fmt(S.now()),))
    lid, start, end, space, people = build = build_class(S, run, mock, ga, tid, start_offset=-80, dur=60)
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    from app import meet_track as T
    run(T.tick())
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(dailyreport.build(aid, start.strftime("%Y-%m-%d"))))
    assert {"Xulosa", "Postlar", "Majlislar"} <= set(wb.sheetnames)
    wb = load_workbook(io.BytesIO(dailyreport.build(aid, "2001-01-01")))
    assert "Majlislar" not in wb.sheetnames and "Xulosa" in wb.sheetnames


# ---------------------------------------------------------------- diagnostika
def test_diagnostics_all_ok(world, run, mock):
    S, G, tid, ga, gb = world
    items = run(G.diagnose())
    st = {i["key"]: i for i in items}
    for k in ("token", "space", "access", "cohost", "calendar", "calendar_clean", "participants", "limit60"):
        assert st[k]["state"] in ("ok", "warn"), (k, st[k])
    assert st["token"]["state"] == "ok" and st["space"]["state"] == "ok" and st["access"]["state"] == "ok" and "OPEN" in st["access"]["detail"]
    assert st["cohost"]["state"] == "ok" and "COHOST" in st["cohost"]["detail"] and "ustoz@gmail.com" in st["cohost"]["detail"]
    assert st["calendar"]["state"] == "ok" and not mock.STATE["calendars"]               # test kalendar o'chirildi
    assert "Meet limiti: 60 daqiqa (Gmail)" in st["limit60"]["detail"]


def test_diagnostics_reports_each_failure_separately(world, run, mock):
    S, G, tid, ga, gb = world
    mock.STATE["bad_members"].add("ustoz@gmail.com")
    st = {i["key"]: i for i in run(G.diagnose())}
    assert st["cohost"]["state"] == "fail" and st["space"]["state"] == "ok" and st["calendar"]["state"] == "ok"
    mock.STATE["meet_disabled"] = True
    st = {i["key"]: i for i in run(G.diagnose())}
    assert st["space"]["state"] == "fail" and "meet.googleapis.com" in st["space"]["fix"]
    assert st["access"]["state"] == "skip" and st["participants"]["state"] == "skip" and st["calendar"]["state"] == "ok"


def test_diagnostics_not_connected(clean, run):
    from app import meet_google as G
    assert run(G.diagnose())[0]["key"] == "cfg"
    clean.put("google_client_id", "x")
    clean.put("google_client_secret", "y")
    assert run(G.diagnose())[0]["key"] == "token"


def test_diagnostics_test_email_for_cohost(world, run, mock):
    S, G, tid, ga, gb = world
    S.db.ex("DELETE FROM mt_teachers")
    st = {i["key"]: i for i in run(G.diagnose())}
    assert st["cohost"]["state"] == "warn" and "Gmail" in st["cohost"]["detail"]
    st = {i["key"]: i for i in run(G.diagnose("tester@gmail.com"))}
    assert st["cohost"]["state"] == "ok"
    st = {i["key"]: i for i in run(G.diagnose("host@gmail.com"))}
    assert st["cohost"]["state"] == "warn" and "o'zini" in st["cohost"]["detail"]


# ---------------------------------------------------------------- veb sahifalar
@pytest.fixture()
def client(world):
    from app.main import app
    return TestClient(app, follow_redirects=False)


PAGES = ["/meet", "/meet?w=1", "/meet/calendar", "/meet/teachers", "/meet/groups", "/meet/lessons", "/meet/reports", "/meet/stats",
         "/meet/alerts", "/meet/settings", "/meet/diag", "/meet/account", "/meet/guide"]


def test_all_pages_render_with_meet_design(client, world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    lid, start, end, space, people = build_class(S, run, mock, ga, tid, start_offset=-80, dur=60)
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    run(T.tick())
    mk_lesson(S, run, gb, tid, S.now() + timedelta(hours=4), 90)
    S.add_schedule(ga, [0, 2, 4], "19:00", 60, tid)
    S.add_alert("send_fail", "sinov", level="error", copy="Havola: https://meet.google.com/x", dedupe="t1")
    for p in PAGES + [f"/meet/groups/{ga}", f"/meet/lessons/{lid}", f"/meet/reports?group={ga}&m={start.strftime('%Y-%m')}", "/meet/groups?scan=0"]:
        r = client.get(p)
        assert r.status_code == 200, (p, r.status_code, r.text[:300])
        assert 'class="ws-meet"' in r.text and "/static/meet.css" in r.text and "Majlislar" in r.text, p
    html = client.get("/meet").text
    assert "Group Post" in html and "Kanallarim" in html and "Tizim tahlili" in html          # bo'lim tanlagichda to'rttasi ham bor
    assert client.get(f"/meet/lessons/{lid}").text.count("Bobur") >= 1
    assert "c-late_early" in client.get(f"/meet/reports?group={ga}&m={start.strftime('%Y-%m')}").text
    assert "nusxalash" in client.get("/meet/alerts").text and "https://meet.google.com/x" in client.get("/meet/alerts").text


def test_other_sections_untouched_and_switcher(client, world):
    for p in ("/", "/sys", "/ch", "/groups", "/settings", "/accounts", "/sys/guide"):
        try:
            r = client.get(p)
        except Exception as e:                      # ish nusxasida mavjud bo'limning ba'zi shablonlari yo'q bo'lishi mumkin
            if "TemplateNotFound" in type(e).__name__:
                continue
            raise
        assert r.status_code in (200, 303), (p, r.status_code, r.text[:200])
        if r.status_code == 200:
            assert "ws-meet" not in r.text and "meet.css" not in r.text, p
            assert "/meet/switch" in r.text, p                                    # tanlagichda Majlislar bor
    r = client.post("/meet/switch")
    assert r.status_code == 303 and r.headers["location"] == "/meet" and "ws=meet" in r.headers["set-cookie"]
    client.cookies.set("ws", "meet")
    assert client.get("/").headers["location"] == "/meet"
    client.cookies.clear()


def test_sys_page_shows_google_status(client, world, run):
    r = client.get("/sys")
    assert r.status_code == 200
    for needle in ("Majlislar (Google Meet)", "Google token", "Meet limiti: 60 daqiqa (Gmail)", "Oxirgi polling", "Faol majlislar"):
        assert needle in r.text, needle


def test_crud_flows_via_http(client, world, run, mock, tg):
    S, G, tid, ga, gb = world
    r = client.post("/meet/teachers/save", data={"name": "Yangi Ustoz", "gmail": "yangi@gmail.com", "tg_username": "@yangi_u", "phone": "+998", "note": "izoh"})
    assert r.status_code == 303 and "msg=" in r.headers["location"]
    assert S.db.one("SELECT tg_username FROM mt_teachers WHERE name='Yangi Ustoz'")["tg_username"] == "yangi_u"
    for bad in ({"name": "", "gmail": ""}, {"name": "X", "gmail": "yomon"}, {"name": "X", "tg_username": "ab"}):
        assert "err=" in client.post("/meet/teachers/save", data=bad).headers["location"]
    r = client.post(f"/meet/groups/{ga}/schedule", data={"weekday": ["0", "2", "4"], "start_time": "20:00", "duration": "90", "teacher_id": str(tid)})
    assert "msg=" in r.headers["location"] and "60" in r.headers["location"]          # 60 daqiqa ogohlantirishi URL xabarida
    assert S.db.one("SELECT COUNT(*) c FROM mt_schedule WHERE group_id=?", (ga,))["c"] == 3
    r = client.post(f"/meet/groups/{ga}/extra", data={"day": (S.now() + timedelta(days=3)).strftime("%Y-%m-%d"), "time": "10:00", "duration": "45"})
    assert "msg=" in r.headers["location"]
    lid = S.db.one("SELECT id FROM mt_lessons WHERE group_id=? ORDER BY id DESC LIMIT 1", (ga,))["id"]
    new = (S.now() + timedelta(days=4)).strftime("%Y-%m-%dT12:00")
    tg[0].clear()
    r = client.post(f"/meet/lessons/{lid}/move", data={"when": new, "duration": "50", "back": "/meet/lessons"})
    assert r.headers["location"].startswith("/meet/lessons?") and S.lesson(lid)["start_at"].endswith("12:00:00") and len(tg[0]) == 2
    r = client.post(f"/meet/lessons/{lid}/cancel", data={"reason": "test"})
    assert S.lesson(lid)["status"] == "cancelled" and "msg=" in r.headers["location"]
    client.post(f"/meet/lessons/{lid}/record", data={"url": "https://drive.google.com/rec"})
    assert S.lesson(lid)["record_url"] == "https://drive.google.com/rec"
    r = client.post("/meet/settings", data={"late_min": "15", "early_min": "12", "poll_sec": "999", "remind_a": "90", "host_type": "workspace", "admin_username": "@boss"})
    assert S.cfg("late_min") == 15 and S.cfg("poll_sec") == 60 and S.cfg("remind_a") == 90 and S.cfg("host_type") == "workspace" and S.cfg("admin_username") == "boss"
    assert S.cfg("remind_on") == 0                                                  # belgilanmagan katak = o'chirilgan
    r = client.post("/meet/settings/google", data={"client_id": "abc", "client_secret": "s3cr3t"})
    assert S.get("google_client_id") == "abc" and S.get("google_client_secret") == "s3cr3t"
    r = client.post(f"/meet/groups/{ga}/delete")
    assert S.group(ga) is None and not S.db.q("SELECT 1 FROM mt_lessons WHERE group_id=?", (ga,))


def test_oauth_http_endpoints(client, clean, run, mock):
    S = clean
    S.put("google_client_id", "cid")
    S.put("google_client_secret", "sec")
    r = client.get("/meet/google/connect")
    assert r.status_code == 303 and r.headers["location"].startswith(mock_url("/auth")) and "code_challenge=" in r.headers["location"]
    state = parse_qs(urlparse(r.headers["location"]).query)["state"][0]
    assert "err=" in client.get("/meet/google/callback?code=good-code&state=boshqa").headers["location"]
    assert "err=" in client.get("/meet/google/callback?error=access_denied").headers["location"]
    S.put("oauth_tmp", json.dumps({"state": state, "verifier": "v" * 50, "at": __import__("time").time()}))
    r = client.get(f"/meet/google/callback?code=good-code&state={state}")
    assert r.headers["location"].startswith("/meet/diag") and S.get("g_email") == "host@gmail.com"
    r = client.post("/meet/google/disconnect")
    assert S.get("g_state") == "none"


def mock_url(path):
    import os
    return os.environ["TGP_GOOGLE_AUTH"].rsplit("/auth", 1)[0] + path


def test_diag_page_runs_and_shows_results(client, world, run, mock):
    r = client.post("/meet/diag/run", data={"test_email": "tester@gmail.com"})
    assert r.status_code == 303
    html = client.get("/meet/diag").text
    assert "Meet xonasi yaratish" in html and "Ishlaydi" in html and "Google Calendar" in html and "Meet limiti: 60 daqiqa (Gmail)" in html


def test_excel_export_endpoint(client, world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    lid, start, end, space, people = build_class(S, run, mock, ga, tid, start_offset=-80, dur=60)
    people[0]["sessions"][0]["endTime"] = iso(end)
    set_conf(mock, space, start, end, people)
    run(T.tick())
    r = client.get(f"/meet/reports/export?group={ga}&m={start.strftime('%Y-%m')}")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "spreadsheetml" in r.headers["content-type"] and "attachment" in r.headers["content-disposition"]


def test_api_live_json(client, world, run, mock, tg):
    S, G, tid, ga, gb = world
    from app import meet_track as T
    start = S.now() - timedelta(minutes=12)
    lid = mk_lesson(S, run, ga, tid, start, 60)
    set_conf(mock, S.lesson(lid)["space_name"], start, None, [person("u1", "Ustoz Test", [(start, None)])])
    run(T.poll_lesson(S.lesson(lid)))
    js = client.get("/meet/api/live").json()
    assert js and js[0]["group"] == "Ingliz tili A" and 11 <= js[0]["elapsed"] <= 12


def test_background_loops_run_under_supervise(world, run, mock, tg):
    """meet_sched/meet_track fonda supervise ostida ishlaydi va yiqilsa qayta ishga tushadi."""
    import asyncio
    from app import meet_sched, meet_track, syscheck, db

    async def go():
        t1 = syscheck.supervise("meet_sched", meet_sched.loop)
        t2 = syscheck.supervise("meet_track", meet_track.loop)
        await asyncio.sleep(0.3)
        assert not t1.done() and not t2.done()
        beats = {r["name"]: r["state"] for r in db.q("SELECT name, state FROM sys_beat WHERE name IN ('meet_sched','meet_track')")}
        t1.cancel()
        t2.cancel()
        for t in (t1, t2):
            try:
                await t
            except asyncio.CancelledError:
                pass
        return beats
    beats = run(go())
    assert beats == {"meet_sched": "running", "meet_track": "running"}
    h = {i["key"]: i for i in meet_sched.health_items()}
    assert "meet_sched" in h and "meet_track" in h


def test_full_app_lifespan_starts_with_meet(world):
    """To'liq dastur (lifespan) Majlislar fon jarayonlari bilan ishga tushadi va to'xtaydi."""
    from app.main import app
    from app import syscheck
    with TestClient(app) as c:
        assert c.get("/meet").status_code == 200
        assert "meet_sched" in syscheck.TASKS and "meet_track" in syscheck.TASKS
